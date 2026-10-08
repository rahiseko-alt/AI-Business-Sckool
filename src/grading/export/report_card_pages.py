"""前期の成績表（Excel）に、印刷用の通知表・成績証明書のシートを式で足す（ADR 0002）。

読むのは「通知表データ」シートだけ（1人1行。4〜63行目。列の決まりは report_card_data.py と同じ）:
A 学籍番号, B 氏名（英字）, C カタカナ, D 学科, E〜CD 科目13枠×6列（科目・形態・点数・評定・設定単位・取得単位。
受けている科目を上に詰め、使わない枠は ""）, CE 前期設定計, CF 前期取得計, CG 授業時数, CH 出席時数,
CI 出席率（0〜1）, CJ 不合格（TRUE/FALSE）, CK 要確認の件数, CL 特記事項。どのセルにも「要確認」が入りうる。

作るシート:
- 通知表_全員: 1人1ページ（A4縦）を60ページ縦に並べ、ページの間に改ページを入れる。k ページ目は通知表データの 3+k 行目。
  どのページも、読む行番号を除けば同じ式。
- 通知表_1人: 黄色のセルに学籍番号を入れると、その学生の1ページになる。無い番号は「要確認」（赤）。
- 成績証明書_1人: 通知表_1人と同じ中身で、題が「成績証明書」。黒塗り・不合格の印・「成績証明書ではありません」の注記は無い。

見た目は、実際に印刷した通知表（report_card.py の render_html の PDF）に合わせた（利用者の指示、2026-10-08）:
学生欄は3行（学科・日本語氏名・学籍番号／学年・英語氏名・発行日／対象期間）、点数と出席時数は整数ならそのまま・
そうでなければ小数第1位まで、E の科目の取得は「－」、前期・後期とも12行、題は1字ずつ空ける。
黒塗り（評定が A〜D・空欄・要確認のどれでもない行）と不合格の印・要確認の赤は条件付き書式なので、
通知表データの値が変われば印刷にもそのまま出る。前期の科目が12を超える学生がいれば、題の左に赤で要確認を出す。

1ページに収まることの前提（ページの寸法は ROW_HEIGHTS と COL_MM が持つ。テストで和を確かめる）:
- 横は1ページ幅に合わせ、縦は手で入れた改ページで区切る。
- 用紙 A4（210×297mm）、余白 左右・上下 0.4インチ（10.2mm）、ヘッダー・フッター 0.2インチ。
  印刷できる幅 189.6mm に対して表の幅は 186mm（印刷した通知表と同じ）。
- 列幅の単位は前期の成績表の既定のフォント（ＭＳ Ｐゴシック 11、数字1文字 8px）で換算する。
"""

from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break

DATA_SHEET = "通知表データ"
ALL_SHEET = "通知表_全員"
ONE_SHEET = "通知表_1人"
CERT_SHEET = "成績証明書_1人"
FIRST = 4            # 通知表データの1人目の行
STUDENTS = 60
SLOTS = 13           # 通知表データの前期の科目の枠
PAGE_SLOTS = 12      # 印刷する前期の行（印刷した通知表と同じ。13枠目に科目があれば要確認）
LATER_ROWS = 12      # 後期の空欄の行（PDF と同じ）
FIELDS = 6           # 科目・形態・点数・評定・設定・取得
TOTALS = {"設定計": "CE", "取得計": "CF", "授業時数": "CG", "出席時数": "CH", "出席率": "CI",
          "不合格": "CJ", "要確認": "CK", "特記事項": "CL"}

# 列 A〜O（mm）。A〜G が1年、H はすきま、I〜O が2年。授業科目は A:B（I:J）を結合
COL_MM = (17, 24, 12, 10, 10, 9, 9, 4, 17, 24, 12, 10, 10, 9, 9)
LAST_COL = "O"
SIDE2 = 9            # 2年の表の左端の列（I）
MARGIN_IN = 0.4
HEADER_IN = 0.2

# 1ページの中の行（ページの先頭からの位置）
R_SCHOOL1, R_SCHOOL2, R_TITLE = 0, 1, 3
R_PROFILE = 5                      # 3行（学科・日本語氏名・学籍番号／学年・英語氏名・発行日／対象期間）
PROFILE_ROWS = 3
R_ATT_H2 = R_PROFILE + PROFILE_ROWS + 1
R_ATT_YEAR, R_ATT_HEAD, R_ATT = R_ATT_H2 + 1, R_ATT_H2 + 2, R_ATT_H2 + 3   # 出席は 前期・後期・年間合計 の3行
R_GR_H2 = R_ATT + 4
R_GR_YEAR, R_GR_HEAD, R_TERM1 = R_GR_H2 + 1, R_GR_H2 + 2, R_GR_H2 + 3
R_SLOT = R_TERM1 + 1               # 12行
R_SUM1 = R_SLOT + PAGE_SLOTS       # 前期計
R_TERM2 = R_SUM1 + 1               # 後期
R_LATER = R_TERM2 + 1              # 12行
R_TOTAL = R_LATER + LATER_ROWS     # 合計取得単位数
R_REM_H2 = R_TOTAL + 2
R_REM = R_REM_H2 + 1               # 3行（結合）
R_FOOT = R_REM + 4                 # 2行
PAGE_ROWS = R_FOOT + 2

_heights = {R_SCHOOL1: 15, R_SCHOOL2: 18, R_TITLE - 1: 4, R_TITLE: 32, R_PROFILE - 1: 4, R_ATT_H2 - 1: 6,
            R_ATT_H2: 16, R_ATT_YEAR: 14, R_GR_H2 - 1: 6, R_GR_H2: 16, R_GR_YEAR: 14, R_TOTAL + 1: 6,
            R_REM_H2: 16, R_REM + 3: 10, R_FOOT: 15, R_FOOT + 1: 15}
ROW_HEIGHTS = tuple(_heights.get(i, 14) for i in range(PAGE_ROWS))   # 書いていない行は表の1行（14pt＝4.9mm）

FONT = "游ゴシック"
# 色は不透明（先頭 FF）で書く。先頭が 00 だと LibreOffice が条件付き書式の塗りを透明として扱い、黒塗りが出ない
THIN = Side(style="thin", color="FF000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WHITE2 = Side(style="double", color="FFFFFFFF")   # 不合格の印の内側の白い二重線
HEAD_FILL = PatternFill("solid", fgColor="FFE8E8E8")
TERM_FILL = PatternFill("solid", fgColor="FFF4F4F4")
YELLOW = PatternFill("solid", fgColor="FFFFFF00")
BLACK = PatternFill("solid", fgColor="FF000000", bgColor="FF000000")
RED = "FFC00000"
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")
DATE = 'yyyy"年"m"月"d"日"'


def slot_col(j: int, field: int) -> str:
    """j 番目（1始まり）の科目の枠の field 番目（0=科目 … 5=取得）の列。"""
    return get_column_letter(5 + FIELDS * (j - 1) + field)


def page_height_pt() -> float:
    """印刷できる高さ（pt）。"""
    return (297 / 25.4 - 2 * MARGIN_IN) * 72


def _width(mm: float) -> float:
    """mm を Excel の列幅（既定のフォントの数字の文字数）に直す。"""
    return round((mm / 25.4 * 96 - 5) / 8, 2)


# ---------- 読む先（全員＝行番号固定、1人＝学籍番号で探す） ----------

class _Row:
    """通知表データの決まった行を読む（通知表_全員の k ページ目）。"""

    def __init__(self, row: int):
        self.row = row

    def ref(self, col: str) -> str:
        return f"'{DATA_SHEET}'!${col}${self.row}"

    def guard(self, expr: str, ident: bool = False) -> str:
        return expr


class _Lookup:
    """学籍番号の入力セルから行を探して読む（通知表_1人・成績証明書_1人）。key は MATCH の結果のセル。"""

    def __init__(self, entry: str, key: str):
        self.entry, self.key = entry, key

    def ref(self, col: str) -> str:
        return f"INDEX('{DATA_SHEET}'!${col}${FIRST}:${col}${FIRST + STUDENTS - 1},{self.key})"

    def guard(self, expr: str, ident: bool = False) -> str:
        # 見つからない番号: 学籍番号・氏名などの欄は「要確認」、ほかの欄は空欄
        missing = f'IF({self.entry}="","","要確認")' if ident else '""'
        return f'IF({self.key}="",{missing},{expr})'


def _value(src, col: str, ident: bool = False) -> str:
    r = src.ref(col)
    return "=" + src.guard(f'IF({r}="","",{r})', ident)


def _number(src, col: str) -> str:
    """数は小数第1位に丸める（整数ならそのまま 48、そうでなければ 72.7。印刷した通知表と同じ）。文字はそのまま。"""
    r = src.ref(col)
    return "=" + src.guard(f'IF({r}="","",IF(ISNUMBER({r}),ROUND({r},1),{r}))')


def _earned(src, j: int) -> str:
    """取得単位。E の科目（取得なし）は「－」、例外処置（評定が F など）は空欄。"""
    r, g = src.ref(slot_col(j, 5)), src.ref(slot_col(j, 3))
    return "=" + src.guard(f'IF({r}="",IF({g}="E","－",""),{r})')


def _date_text() -> str:
    """今日の日付を「2026年10月8日」の文字で（Excel・LibreOffice の言語設定に左右されない）。"""
    return 'YEAR(TODAY())&"年"&MONTH(TODAY())&"月"&DAY(TODAY())&"日"'


# ---------- 書式 ----------

def _put(ws, row: int, col: int, value=None, *, last_col: int | None = None, size: float = 9, bold=False,
         align=LEFT, border=False, fill=None, fmt=None, color=None, wrap=False, shrink=False):
    """セルに値を入れて書式を付ける。last_col があれば col〜last_col を結合する。"""
    last = last_col or col
    if last != col:
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=last)
    cell = ws.cell(row, col)
    if value is not None:
        cell.value = value
    for c in range(col, last + 1):
        x = ws.cell(row, c)
        x.font = Font(name=FONT, size=size, bold=bold, color=color)
        x.alignment = Alignment(horizontal=align.horizontal, vertical="center", wrap_text=wrap, shrink_to_fit=shrink)
        if border:
            x.border = BOX
        if fill:
            x.fill = fill
        if fmt:
            x.number_format = fmt
    return cell


def _head(ws, row: int, col: int, text: str, last_col: int | None = None):
    _put(ws, row, col, text, last_col=last_col, size=9, align=CENTER, border=True, fill=HEAD_FILL, shrink=True)


def _h2(ws, row: int, text: str):
    _put(ws, row, 1, text, last_col=len(COL_MM), size=10.5, bold=True)
    ws.cell(row, 1).border = Border(left=Side(style="thick", color="FF000000"))


# ---------- 1ページ ----------

def _attendance(ws, top: int, side: int, src=None):
    """出席の表。side: 1＝1年（A〜G）、SIDE2＝2年（I〜O）。src が無ければ見出しだけ。
    列: 区分（a）・授業時数（a+1）・出席時数（a+2:a+3）・出席率（a+4:a+6）。"""
    a = side
    year = "1年（2026年度）" if src else "2年（2027年度）"
    _put(ws, top + R_ATT_YEAR, a, year, last_col=a + 6, bold=True)
    _head(ws, top + R_ATT_HEAD, a, "区分")
    _head(ws, top + R_ATT_HEAD, a + 1, "授業時数")
    _head(ws, top + R_ATT_HEAD, a + 2, "出席時数", a + 3)
    _head(ws, top + R_ATT_HEAD, a + 4, "出席率", a + 6)
    for i, label in enumerate(("前期", "後期", "年間合計")):
        r = top + R_ATT + i
        filled = src is not None and label != "後期"   # 年間合計は前期と同じ（後期はまだ無い）
        bold = label == "年間合計"
        _put(ws, r, a, label, border=True, fill=HEAD_FILL, align=CENTER, bold=bold)
        _put(ws, r, a + 1, _value(src, TOTALS["授業時数"]) if filled else None, border=True, align=RIGHT, bold=bold)
        _put(ws, r, a + 2, _number(src, TOTALS["出席時数"]) if filled else None, last_col=a + 3, border=True,
             align=RIGHT, bold=bold)
        _put(ws, r, a + 4, _value(src, TOTALS["出席率"]) if filled else None, last_col=a + 6, border=True,
             align=RIGHT, fmt="0.0%", bold=bold)


def _grades(ws, top: int, side: int, src=None):
    """成績・単位の表。前期は受けている科目を上から詰めて12行、前期計、後期（空欄）、合計取得単位数。
    列: 授業科目（a:a+1）・点数・評定・形態・設定・取得（a+2〜a+6）。"""
    a = side
    _put(ws, top + R_GR_YEAR, a, "1年" if src else "2年", last_col=a + 6, bold=True)
    _head(ws, top + R_GR_HEAD, a, "授業科目", a + 1)
    for k, h in enumerate(("点数", "評定", "形態", "設定", "取得"), 2):
        _head(ws, top + R_GR_HEAD, a + k, h)
    for r, label in ((top + R_TERM1, "前期"), (top + R_TERM2, "後期")):
        _put(ws, r, a, label, last_col=a + 6, bold=True, border=True, fill=TERM_FILL)
    for j in range(1, PAGE_SLOTS + 1):
        r = top + R_SLOT + j - 1
        cells = ((_value(src, slot_col(j, 0)), LEFT), (_number(src, slot_col(j, 2)), RIGHT),
                 (_value(src, slot_col(j, 3)), CENTER), (_value(src, slot_col(j, 1)), CENTER),
                 (_value(src, slot_col(j, 4)), RIGHT), (_earned(src, j), RIGHT)) if src else ((None, LEFT),) * 6
        _put(ws, r, a, cells[0][0], last_col=a + 1, border=True, align=cells[0][1], size=8.5, shrink=True)
        for k, (v, align) in enumerate(cells[1:], 2):
            _put(ws, r, a + k, v, border=True, align=align, shrink=True)
    for i in range(LATER_ROWS):
        _put(ws, top + R_LATER + i, a, last_col=a + 1, border=True)
        for k in range(2, 7):
            _put(ws, top + R_LATER + i, a + k, border=True)
    for r, label in ((top + R_SUM1, "前期計"), (top + R_TOTAL, f"{'1年' if src else '2年'} 合計取得単位数")):
        last = a + 1 if r == top + R_SUM1 else a + 4
        _put(ws, r, a, label, last_col=last, bold=True, border=True)
        for k in range(last - a + 1, 5):
            _put(ws, r, a + k, border=True)
        _put(ws, r, a + 5, _value(src, TOTALS["設定計"]) if src else None, border=True, align=RIGHT, bold=True)
        _put(ws, r, a + 6, _value(src, TOTALS["取得計"]) if src else None, border=True, align=RIGHT, bold=True)


def _page(ws, top: int, src, kind: str, note: str):
    """top 行目から1ページ。kind: "通知表" か "成績証明書"。note: 題の左の赤い注意の式。"""
    report = kind == "通知表"
    end = len(COL_MM)
    _put(ws, top + R_SCHOOL1, 1, "学校法人海鵬学園", last_col=7, size=10)
    _put(ws, top + R_SCHOOL2, 1, "AIビジネス専門学校", last_col=7, size=12, bold=True)
    if report:
        _put(ws, top + R_SCHOOL1, 11, "この書類は成績証明書ではありません", last_col=end, align=CENTER, border=True,
             shrink=True)
    t = top + R_TITLE
    _put(ws, t, 1, note, last_col=2, size=10, bold=True, color=RED, wrap=True)
    _put(ws, t, 3, "　".join(kind), last_col=11, size=18, bold=True, align=CENTER)
    if report:
        stamp = src.ref(TOTALS["不合格"])
        _put(ws, t, 13, "=" + src.guard(f'IF({stamp}=TRUE,"不合格","")'), last_col=end, size=18, bold=True,
             align=CENTER, shrink=True)
    p = top + R_PROFILE
    rows = (("学科", _value(src, "D", True), "日本語氏名", _value(src, "C", True), "学籍番号", _value(src, "A", True)),
            ("学年", "1年", "英語氏名", _value(src, "B", True), "発行日", "=TODAY()"))
    for i, (l1, v1, l2, v2, l3, v3) in enumerate(rows):
        _head(ws, p + i, 1, l1)
        _put(ws, p + i, 2, v1, border=True, shrink=True)
        _head(ws, p + i, 3, l2, 4)
        _put(ws, p + i, 5, v2, last_col=10, border=True, shrink=True)   # 長い英語の氏名は縮めて1行に収める
        _head(ws, p + i, 11, l3, 12)
        _put(ws, p + i, 13, v3, last_col=end, border=True, shrink=True, fmt=DATE if v3 == "=TODAY()" else None)
    _head(ws, p + 2, 1, "対象期間")
    _put(ws, p + 2, 2, "2026年度 前期", last_col=end, border=True)
    _h2(ws, top + R_ATT_H2, "出席（1年・2年／前期・後期）")
    _attendance(ws, top, 1, src)
    _attendance(ws, top, SIDE2)
    _h2(ws, top + R_GR_H2, "成績・単位")
    _grades(ws, top, 1, src)
    _grades(ws, top, SIDE2)
    _h2(ws, top + R_REM_H2, "特記事項")
    rem = src.ref(TOTALS["特記事項"])
    ws.merge_cells(start_row=top + R_REM, start_column=1, end_row=top + R_REM + 2, end_column=end)
    ws.cell(top + R_REM, 1).value = "=" + src.guard(f'IF({rem}="","該当なし",{rem})')
    for r in range(top + R_REM, top + R_REM + 3):
        for c in range(1, end + 1):
            ws.cell(r, c).border = BOX
            ws.cell(r, c).font = Font(name=FONT, size=9)
            ws.cell(r, c).alignment = Alignment(vertical="top", wrap_text=True)
    f = top + R_FOOT
    _put(ws, f, 1, "記載の期間における出席及び成績は、上記のとおりです。", last_col=7)
    _put(ws, f, 9, f'="発行年月日　"&{_date_text()}', last_col=end, align=RIGHT)
    _put(ws, f + 1, 9, "学校法人海鵬学園　AIビジネス専門学校", last_col=end, align=RIGHT)
    for i, h in enumerate(ROW_HEIGHTS):
        ws.row_dimensions[top + i].height = h
    _rules(ws, top, report)


def _rules(ws, top: int, report: bool):
    """条件付き書式: 「要確認」は赤。通知表は E・例外の行を黒地に白文字、不合格の印を黒地に白文字。"""
    bottom = top + PAGE_ROWS - 1
    ws.conditional_formatting.add(f"A{top}:{LAST_COL}{bottom}", CellIsRule(
        operator="equal", formula=['"要確認"'], font=Font(color=RED, bold=True)))
    if not report:
        return
    s1, s2 = top + R_SLOT, top + R_SLOT + PAGE_SLOTS - 1
    g = f"$D{s1}"
    black = f'AND({g}<>"",{g}<>"A",{g}<>"B",{g}<>"C",{g}<>"D",{g}<>"要確認")'
    ws.conditional_formatting.add(f"A{s1}:G{s2}", FormulaRule(
        formula=[black], fill=BLACK, font=Font(color="FFFFFFFF")))
    t = top + R_TITLE
    ws.conditional_formatting.add(f"M{t}:O{t}", CellIsRule(
        operator="equal", formula=['"不合格"'], fill=BLACK, font=Font(color="FFFFFFFF", bold=True),
        border=Border(left=WHITE2, right=WHITE2, top=WHITE2, bottom=WHITE2)))


def _note(src) -> str:
    """題の左の赤い注意。13枠目に科目がある（印刷の12行に収まらない）か、要確認の件数が1以上
    （または件数そのものが読めない）なら「要確認あり」。"""
    n = src.ref(TOTALS["要確認"])
    over = src.ref(slot_col(SLOTS, 0))
    return "=" + src.guard(f'IF({over}<>"","要確認：前期の科目が{PAGE_SLOTS}を超えています",'
                           f'IF(ISNUMBER({n}),IF({n}>0,"要確認あり（"&{n}&"件）",""),IF({n}="","","要確認あり")))')


def _sheet(wb, title: str):
    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title)
    for i, mm in enumerate(COL_MM, 1):
        ws.column_dimensions[get_column_letter(i)].width = _width(mm)
    ws.sheet_view.showGridLines = False
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    # 横は1ページ幅に合わせる（フォントの違いで右端がはみ出して2ページに割れるのを防ぐ）。
    # 縦は「自動」（0）にして、手で入れた改ページで1人1ページに区切る
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    m = ws.page_margins
    m.left = m.right = m.top = m.bottom = MARGIN_IN
    m.header = m.footer = HEADER_IN
    ws.print_options.horizontalCentered = True
    return ws


def _single(wb, title: str, kind: str):
    """1人ぶん。1行目の黄色のセル（B1）に学籍番号を入れる。Q1 が通知表データの何人目か（見つからなければ空）。"""
    ws = _sheet(wb, title)
    _put(ws, 1, 1, "学籍番号を入力 →", bold=True, align=RIGHT)
    _put(ws, 1, 2, None, last_col=4, fill=YELLOW, border=True, size=11)
    ws.cell(1, 2).number_format = "@"
    rng = f"'{DATA_SHEET}'!$A${FIRST}:$A${FIRST + STUDENTS - 1}"
    _put(ws, 1, 16, "何人目", align=RIGHT, size=8)
    # 文字で入っていても数で入っていても見つける
    ws["Q1"] = (f'=IF($B$1="","",IFERROR(MATCH($B$1,{rng},0),IFERROR(MATCH($B$1&"",{rng},0),'
                f'IFERROR(MATCH(VALUE($B$1),{rng},0),""))))')
    src = _Lookup("$B$1", "$Q$1")
    note = ('=IF($B$1="","学籍番号を黄色の欄に入れてください",IF($Q$1="","要確認：学籍番号が見つかりません",'
            + _note(src)[1:] + "))")
    top = 3
    _page(ws, top, src, kind, note)
    ws.print_area = f"A{top}:{LAST_COL}{top + PAGE_ROWS - 1}"
    return ws


def add_pages(wb, students: int = STUDENTS):
    """通知表_全員・通知表_1人・成績証明書_1人を wb に足す（同じ名前のシートがあれば作り直す）。"""
    ws = _sheet(wb, ALL_SHEET)
    for k in range(1, students + 1):
        top = 1 + (k - 1) * PAGE_ROWS
        src = _Row(FIRST + k - 1)
        _page(ws, top, src, "通知表", _note(src))
        if k < students:
            ws.row_breaks.append(Break(id=top + PAGE_ROWS - 1))
    ws.print_area = f"A1:{LAST_COL}{students * PAGE_ROWS}"
    _single(wb, ONE_SHEET, "通知表")
    _single(wb, CERT_SHEET, "成績証明書")
    return wb
