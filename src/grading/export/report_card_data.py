"""通知表のデータを、成績表の Excel の中に式で作る（ADR 0002）。

成績表のブック（式のまま読んだもの）に次のシートを足す。どのセルも AI計算版と出席簿の月別シートを参照する式で、
点数・評定・出欠の数字は1つも書き込まない。成績表や出席簿を直せば、再計算で全部が変わる。

- 「通知表データ」: 1行1人（国際→総合、AI計算版の行順）。印刷用のページはここだけを見る（下の「契約」）
- 「例外」: 例外処置（点数・取得を空欄、評定を指定の文字にする）。1行1件、理由とともに書く
- 「特記事項」: 学生ごとの特記事項の文。空なら「該当なし」
- 「確認」: 「要確認」の件数（B2）と一覧（学籍番号・科目・何がおかしいか）。0件になるまで印刷しない
- 「通知表_計算」「通知表_授業列」: 途中の計算（学生×科目の点数・評定・出欠、出席簿の列ごとの「数える列か」）

決まりは report_card.py（照合の正）と同じ:
- 受けている科目＝評価項目に数値が1つでもある科目。ひな形（TEMPLATE）の順に、受けている科目だけを上から詰める
- 点数＝AI計算版の合計（丸めない。表示は印刷側で小数第1位）、評定＝AI計算版の評定。取得＝設定（E と例外は空欄）
- 授業時数＝受けている科目の授業数（×を除く）の合計。出席時数＝授業時数−欠席−遅刻÷3（丸めない）
- 授業数・欠席・遅刻は subject_counts と同じく、出席簿の4行目の科目名が一致する列の「×以外」「欠」「遅」を数える。
  日本語運用力強化演習も1コマずつ数える（出席点の「週ごと」とは違う）
- 決められない値（評定が A〜E でない、点数が数値でない、授業数が0）は推測せず、その欄に「要確認」と出す

式で表しきれず、read_register と食い違いうる点（2026年度前期の出席簿では、どれも結果は同じになる）:
- 「数える列か」は、3行目の見出しを左から見て、数字の入った見出し（日付・日付の式を含む）の下なら数える、
  数字の無い「月　日（月）」の下なら数えない、で判定する。read_register の日付の読み取り（曜日の違い・訂正
  DATE_CORRECTIONS）は数えるかどうかを変えないので式には入れていない。ただし、数字はあるが日付として読めない見出し
  は、read_register では数えないが式では数える
- 出席簿の科目名は、作った時点の4行目の書き方（康熙部首・全角・後ろの空白の違いを含む）で照合する。4行目を新しい
  書き方に直した列は数えられなくなる（その科目の授業数が減り、0なら「要確認」）
- 学生が月別シートのどの行にいるかは、作った時点の行で固定する（学生は60人で固定。ADR 0002）
"""

from collections.abc import Sequence

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L

from grading.export.linked import describe
from grading.export.report_card import REGISTER_NAMES, TEMPLATE
from grading.importing.attendance import _end_column, _month_sheets, unify_subject

DEPTS = (  # （AI計算版のシート, 学科名, 出席簿シートの頭）。build_report_cards.py と同じ
    ("AI計算版_国際", "国際ビジネス科", "国際ビジネスAI科"),
    ("AI計算版_総合", "総合ビジネス科", "総合ビジネス科"),
)
# 利用者の指示による例外処置（2026年度前期）。科目は通知表の科目名で書く
EXCEPTIONS_2026_FIRST = (("AIBC26049", "日本語能力強化演習", "F", "利用者の指示による例外処置"),)

DATA, EXCEPTION, REMARK, CHECK = "通知表データ", "例外", "特記事項", "確認"
CALC, COLUMNS = "通知表_計算", "通知表_授業列"
UNSURE = "要確認"
FIRST = 4                         # 通知表データ・通知表_計算の1人目の行
SLOTS = len(TEMPLATE)             # 13
SLOT_FIELDS = ("科目", "形態", "点数", "評定", "設定", "取得")
TAIL = ("前期設定計", "前期取得計", "授業時数", "出席時数", "出席率", "不合格", "要確認の件数", "特記事項")
EXCEPTION_ROWS = 100              # 例外シートの記入欄（2〜101行）
REMARK_ROWS = 200                 # 特記事項シートで探す行（2〜201行）
_HEAD = PatternFill("solid", fgColor="DDE7F3")
_INPUT = PatternFill("solid", fgColor="FFF2CC")
_DIGITS = "{" + ",".join(f'"{d}"' for d in "0123456789０１２３４５６７８９") + "}"

# 通知表_計算 の項目（1項目＝ひな形の13科目ぶんの列）
CALC_FIELDS = ("受講", "順番", "枠", "点数", "評定", "取得", "授業数", "欠席", "遅刻", "例外", "例外の評定")


def slot_column(j: int, field: str) -> int:
    """通知表データで、j 番目（1始まり）の枠の field の列番号。"""
    return 5 + 6 * (j - 1) + SLOT_FIELDS.index(field)


def tail_column(name: str) -> int:
    return 5 + 6 * SLOTS + TAIL.index(name)


def _calc_col(field: str, t: int) -> int:
    """通知表_計算で、field の t 番目（1始まり、ひな形の順）の科目の列。A 列は学籍番号。"""
    return 2 + CALC_FIELDS.index(field) * SLOTS + (t - 1)


def _calc_row_range(field: str, r: int) -> str:
    return f"{_q(CALC)}!${L(_calc_col(field, 1))}${r}:${L(_calc_col(field, SLOTS))}${r}"


def _q(title: str) -> str:
    return "'" + title.replace("'", "''") + "'"


def _text(value: str) -> str:
    return '"' + value.replace('"', '""') + '"'


def _array(values) -> str:
    return "{" + ",".join(_text(v) if isinstance(v, str) else str(v) for v in values) + "}"


def _head(row):
    for cell in row:
        cell.font, cell.fill = Font(bold=True), _HEAD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _student_rows(ws) -> dict[str, int]:
    return {str(ws.cell(r, 2).value).strip(): r for r in range(6, ws.max_row + 1) if ws.cell(r, 2).value}


# ---------- 出席簿: 数える列か ----------

def _columns_sheet(wb, depts) -> dict[str, int]:
    """月別シート1枚につき1行、月別シートと同じ列に「数える列なら1」を式で並べる。返り値: シート名 → 行。"""
    ws = wb.create_sheet(COLUMNS)
    ws["A1"] = ("出席簿の月別シートの各列を数えるか（1＝数える）。3行目の見出しを左から見て、数字の入った見出し（日付）"
                "の下の列は数え、数字の無い「月　日（曜）」の下の列は数えない。列は月別シートと同じ位置")
    ws["A1"].font = Font(bold=True)
    rows, r = {}, 3
    for _, _, prefix in depts:
        for month in _month_sheets(wb, prefix):
            end = _end_column(month)
            ws.cell(r, 1, month.title)
            for c in range(5, end):
                h = f"{_q(month.title)}!{L(c)}$3"
                prev = "0" if c == 5 else f"{L(c - 1)}{r}"
                ws.cell(r, c, f"=IF(ISERROR({h}),0,IF(ISNUMBER({h}),1,IF(TRIM({h}&\"\")=\"\",{prev},"
                              f"IF(SUMPRODUCT(--ISNUMBER(FIND({_DIGITS},{h}&\"\")))>0,1,"
                              f"IF(AND(ISNUMBER(FIND(\"月\",{h}&\"\")),ISNUMBER(FIND(\"日\",{h}&\"\"))),0,{prev})))))")
            rows[month.title] = r
            r += 1
    ws.column_dimensions["A"].width = 24
    return rows


def _register_heads(month, ai_name: str) -> list[str]:
    """この月別シートの4行目で、AI計算版の科目 ai_name にあたる見出し（原文）。"""
    out = []
    for c in range(5, _end_column(month)):
        head = month.cell(4, c).value
        if head is None:
            continue
        unified = unify_subject(str(head))
        if REGISTER_NAMES.get(unified, unified) == unify_subject(ai_name) and str(head) not in out:
            out.append(str(head))
    return out


def _register_subjects(wb, prefix: str, ai_name: str) -> set[str]:
    found = set()
    for month in _month_sheets(wb, prefix):
        for head in _register_heads(month, ai_name):
            found.add(unify_subject(head))
    return found


def _count(wb, prefix: str, flags: dict[str, int], ai_name: str, sid: str, condition: str) -> str:
    """出席簿で ai_name にあたる列のうち、数える列で condition（例 ="欠"）を満たすセルの数。"""
    terms = []
    for month in _month_sheets(wb, prefix):
        heads = _register_heads(month, ai_name)
        rows = _student_rows(month)
        if not heads or sid not in rows:
            continue
        last = L(_end_column(month) - 1)
        head = f"{_q(month.title)}!$E$4:${last}$4"
        match = "+".join(f"({head}={_text(h)})" for h in heads)
        match = f"({match})" if len(heads) == 1 else f"(({match})>0)"
        flag = f"{_q(COLUMNS)}!$E${flags[month.title]}:${last}${flags[month.title]}"
        cells = f"{_q(month.title)}!$E${rows[sid]}:${last}${rows[sid]}"
        terms.append(f"SUMPRODUCT({flag}*{match}*({cells}{condition}))")
    return "=" + ("+".join(terms) if terms else "0")


# ---------- 本体 ----------

def add_report_card_data(wb, depts: Sequence[tuple[str, str, str]] = DEPTS,
                         exceptions: Sequence[tuple[str, str, str, str]] = ()) -> int:
    """成績表のブック（式のまま読んだもの）に通知表データ・例外・特記事項・確認と途中の計算のシートを足す。
    exceptions: 例外シートに最初から書いておく行（学籍番号, 通知表の科目名, 表示する評定, 理由）。返り値: 学生の人数。"""
    for title in (DATA, EXCEPTION, REMARK, CHECK, CALC, COLUMNS):
        if title in wb.sheetnames:
            raise ValueError(f"「{title}」シートがもう有る（例外・特記事項の記入を消さないよう、作り直さない）")
    template_ai = [unify_subject(g) for *_, g in TEMPLATE]
    students = []      # （学科, AI計算版のシート, 行, 学籍番号, 出席簿シートの頭, {t: Subject}）
    for sheet, dept, prefix in depts:
        ws = wb[sheet]
        d = describe(ws, dept, dept, {})
        by_t = {}
        for s in d.subjects:
            if unify_subject(s.name) not in template_ai:
                raise ValueError(f"{dept}: 通知表のひな形に無い科目 {s.name}")
            if len(_register_subjects(wb, prefix, s.name)) > 1:
                raise ValueError(f"{dept}: 出席簿で {s.name} にあたる科目名が2つ以上ある")
            by_t[template_ai.index(unify_subject(s.name)) + 1] = s
        for r in d.rows:
            students.append((dept, sheet, r, str(ws.cell(r, 2).value).strip(), prefix, by_t))
    last = FIRST + len(students) - 1

    flags = _columns_sheet(wb, depts)
    data = wb.create_sheet(DATA)
    calc = wb.create_sheet(CALC)
    exc = _exception_sheet(wb, exceptions, last)
    _remark_sheet(wb, last)

    # 通知表_計算: 見出し
    calc["A1"] = "通知表の途中の計算（学生×ひな形の13科目。通知表データはここから受けている科目を詰めて並べる）"
    calc["A1"].font = Font(bold=True)
    calc.cell(3, 1, "学籍番号")
    for f in CALC_FIELDS:
        calc.cell(2, _calc_col(f, 1), f)
        calc.merge_cells(start_row=2, start_column=_calc_col(f, 1), end_row=2, end_column=_calc_col(f, SLOTS))
        for t, (name, *_rest) in enumerate(TEMPLATE, 1):
            calc.cell(3, _calc_col(f, t), name if f != "枠" else f"{t}枠目")
    _head(calc[2])
    _head(calc[3])

    names = _array(n for n, *_ in TEMPLATE)
    kinds = _array(k for _, k, _, _ in TEMPLATE)
    credits = _array(c for _, _, c, _ in TEMPLATE)
    exc_key = f"{_q(EXCEPTION)}!$F$2:$F${1 + EXCEPTION_ROWS}"
    exc_grade = f"{_q(EXCEPTION)}!$C$2:$C${1 + EXCEPTION_ROWS}"

    for i, (dept, sheet, src, sid, prefix, by_t) in enumerate(students):
        r = FIRST + i
        ai = _q(sheet)
        calc.cell(r, 1, f"={_q(DATA)}!$A{r}")

        def at(field, t, rr=r):
            return f"{L(_calc_col(field, t))}{rr}"

        for t, (name, _kind, credit, _g) in enumerate(TEMPLATE, 1):
            s = by_t.get(t)
            if s is None:          # この学科に無い科目
                for f in CALC_FIELDS:
                    calc[at(f, t)] = 0 if f in ("受講", "授業数", "欠席", "遅刻", "例外") else '=""'
                calc[at("枠", t)] = f"=IFERROR(MATCH({t},{_calc_row_range('順番', r)},0),0)"
                continue
            items = ",".join(f"{ai}!${L(c)}${src}" for c in s.items)
            total, grade = f"{ai}!${L(s.total)}${src}", f"{ai}!${L(s.grade)}${src}"
            taken, withheld = at("受講", t), at("例外", t)
            calc[taken] = f"=IF(COUNT({items})>0,1,0)"
            calc[at("順番", t)] = f'=IF({taken}=1,SUM(${L(_calc_col("受講", 1))}{r}:{taken}),"")'
            calc[at("枠", t)] = f"=IFERROR(MATCH({t},{_calc_row_range('順番', r)},0),0)"
            calc[withheld] = f'=IF({taken}=1,IF(COUNTIF({exc_key},$A{r}&"|"&{_text(name)})>0,1,0),0)'
            calc[at("例外の評定", t)] = (f'=IF({withheld}=1,INDEX({exc_grade},MATCH($A{r}&"|"&{_text(name)},{exc_key},0))&"","")')
            calc[at("点数", t)] = f'=IF({taken}<>1,"",IF({withheld}=1,"",IF(ISNUMBER({total}),{total},"{UNSURE}")))'
            ok = ",".join(f'EXACT({grade},"{g}")' for g in "ABCDE")
            calc[at("評定", t)] = (f'=IF({taken}<>1,"",IF({withheld}=1,{at("例外の評定", t)},'
                                   f'IFERROR(IF(OR({ok}),{grade},"{UNSURE}"),"{UNSURE}")))')
            g = at("評定", t)
            calc[at("取得", t)] = (f'=IF({taken}<>1,"",IF({withheld}=1,"",IF({g}="{UNSURE}","{UNSURE}",'
                                   f'IF({g}="E","",{credit}))))')
            calc[at("授業数", t)] = _count(wb, prefix, flags, s.name, sid, '<>"×"')
            calc[at("欠席", t)] = _count(wb, prefix, flags, s.name, sid, '="欠"')
            calc[at("遅刻", t)] = _count(wb, prefix, flags, s.name, sid, '="遅"')

        # 通知表データ
        data.cell(r, 1, f'=TRIM({ai}!$B${src}&"")')
        data.cell(r, 2, f'=TRIM(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({ai}!$C${src}&"",CHAR(10)," "),CHAR(9)," "),"　"," "))')
        data.cell(r, 3, f'=TRIM(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({ai}!$D${src}&"",CHAR(10)," "),CHAR(9)," "),"　"," "))')
        data.cell(r, 4, dept)
        for j in range(1, SLOTS + 1):
            t = f"{_q(CALC)}!{L(_calc_col('枠', j))}{r}"
            put = {
                "科目": f"INDEX({names},1,{t})", "形態": f"INDEX({kinds},1,{t})",
                "点数": f"INDEX({_calc_row_range('点数', r)},1,{t})", "評定": f"INDEX({_calc_row_range('評定', r)},1,{t})",
                "設定": f"INDEX({credits},1,{t})", "取得": f"INDEX({_calc_row_range('取得', r)},1,{t})",
            }
            for field, expr in put.items():
                cell = data.cell(r, slot_column(j, field), f'=IF({t}=0,"",{expr})')
                if field == "点数":
                    cell.number_format = "0.0"
        sets = ",".join(f"{L(slot_column(j, '設定'))}{r}" for j in range(1, SLOTS + 1))
        earns = ",".join(f"{L(slot_column(j, '取得'))}{r}" for j in range(1, SLOTS + 1))
        col = {n: L(tail_column(n)) for n in TAIL}
        taken_row, n_row = _calc_row_range("受講", r), _calc_row_range("授業数", r)
        a_row, late_row = _calc_row_range("欠席", r), _calc_row_range("遅刻", r)
        grade_row, withheld_row = _calc_row_range("評定", r), _calc_row_range("例外", r)
        data[f"{col['前期設定計']}{r}"] = f"=SUM({sets})"
        data[f"{col['前期取得計']}{r}"] = (f'=IF(COUNTIF({_calc_row_range("取得", r)},"{UNSURE}")>0,"{UNSURE}",SUM({earns}))')
        data[f"{col['授業時数']}{r}"] = (f'=IF(OR(SUM({taken_row})=0,SUMPRODUCT({taken_row}*({n_row}=0))>0),'
                                       f'"{UNSURE}",SUMPRODUCT({taken_row},{n_row}))')
        hours = f"{col['授業時数']}{r}"
        data[f"{col['出席時数']}{r}"] = (f'=IF(ISNUMBER({hours}),SUMPRODUCT({taken_row}*({n_row}-{a_row}-{late_row}/3)),'
                                       f'"{UNSURE}")')
        data[f"{col['出席時数']}{r}"].number_format = "0.0"
        data[f"{col['出席率']}{r}"] = f'=IF(ISNUMBER({hours}),{col["出席時数"]}{r}/{hours},"{UNSURE}")'
        data[f"{col['出席率']}{r}"].number_format = "0.0%"
        data[f"{col['不合格']}{r}"] = (f'=IF(SUMPRODUCT(--({grade_row}="E"))+SUM({withheld_row})>0,TRUE,'
                                     f'IF(COUNTIF({grade_row},"{UNSURE}")>0,"{UNSURE}",FALSE))')
        data[f"{col['要確認の件数']}{r}"] = f'=COUNTIF(A{r}:{col["不合格"]}{r},"{UNSURE}")'
        note = (f'IFERROR(INDEX({_q(REMARK)}!$B$2:$B${1 + REMARK_ROWS},'
                f'MATCH($A{r},{_q(REMARK)}!$A$2:$A${1 + REMARK_ROWS},0))&"","")')
        data[f"{col['特記事項']}{r}"] = f'=IF(TRIM(SUBSTITUTE({note},"　",""))="","該当なし",TRIM({note}))'

    # 通知表データ: 見出し
    data["A1"] = "通知表データ（成績表・出席簿・例外・特記事項から式で作る。直すのは元のシートで、ここには書かない）"
    data["A1"].font = Font(bold=True, size=12)
    data["A2"] = f"「{UNSURE}」が1つでもあれば印刷しない（件数は「{CHECK}」シートの B2）"
    heads = ["学籍番号", "氏名", "カタカナ", "学科"]
    for j in range(1, SLOTS + 1):
        heads += [f"{f}{j}" for f in SLOT_FIELDS]
    heads += list(TAIL)
    for c, h in enumerate(heads, 1):
        data.cell(3, c, h)
    _head(data[3])
    data.freeze_panes = "E4"
    data.column_dimensions["A"].width, data.column_dimensions["B"].width = 12, 28
    data.column_dimensions["C"].width, data.column_dimensions["D"].width = 20, 14

    _check_sheet(wb, students, exc, last)
    calc.freeze_panes = "B4"
    return len(students)


def _exception_sheet(wb, exceptions, last: int):
    ws = wb.create_sheet(EXCEPTION)
    for c, h in enumerate(["学籍番号", "科目（通知表の科目名）", "表示する評定", "理由", "", "照合用", "状態"], 1):
        ws.cell(1, c, h or None)
    _head([ws.cell(1, c) for c in (1, 2, 3, 4, 6, 7)])
    for k, row in enumerate(exceptions):
        for c, v in enumerate(row, 1):
            ws.cell(2 + k, c, v)
    names = f"{_q(CALC)}!${L(_calc_col('受講', 1))}$3:${L(_calc_col('受講', SLOTS))}$3"
    ids = f"{_q(DATA)}!$A${FIRST}:$A${last}"
    taken = f"{_q(CALC)}!${L(_calc_col('受講', 1))}${FIRST}:${L(_calc_col('受講', SLOTS))}${last}"
    for r in range(2, 2 + EXCEPTION_ROWS):
        a, b = f'TRIM($A{r}&"")', f'TRIM($B{r}&"")'
        ws.cell(r, 6, f'=IF(AND({a}="",{b}=""),"",{a}&"|"&{b})')
        ws.cell(r, 7, (f'=IF($F{r}="","",IF(OR({a}="",{b}=""),"{UNSURE}：学籍番号と科目の両方を書く",'
                       f'IF(ISNA(MATCH({a},{ids},0)),"{UNSURE}：この学籍番号の学生がいない",'
                       f'IF(ISNA(MATCH({b},{names},0)),"{UNSURE}：通知表の科目名と違う",'
                       f'IF(INDEX({taken},MATCH({a},{ids},0),MATCH({b},{names},0))<>1,"{UNSURE}：この学生はこの科目を受けていない",'
                       f'IF(COUNTIF($F$2:$F${1 + EXCEPTION_ROWS},$F{r})>1,"{UNSURE}：同じ学生・科目の行が2つ以上ある",'
                       '"反映済み"))))))'))
        for c in range(1, 5):
            ws.cell(r, c).fill = _INPUT
    ws["I1"] = ("例外処置: 成績表の点数は変えずに、通知表の上だけで その科目の点数・取得を空欄、評定を「表示する評定」にする。"
                "その学生には不合格印が付く。科目は通知表の科目名（例: ビジネス日本語Ⅰ、日本語能力強化演習）で書く。"
                "黄色の欄に1行1件。「状態」が「要確認」の行は通知表に反映されない")
    ws["I1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["I"].width = 60
    ws.column_dimensions["F"].hidden = True
    for c, w in zip("ABCDG", (12, 28, 10, 32, 40)):
        ws.column_dimensions[c].width = w
    ws.freeze_panes = "A2"
    return ws


def _remark_sheet(wb, last: int):
    ws = wb.create_sheet(REMARK)
    for c, h in enumerate(["学籍番号", "特記事項（空なら「該当なし」）", "氏名（参考）"], 1):
        ws.cell(1, c, h)
    _head(ws[1])
    for k, r in enumerate(range(FIRST, last + 1)):
        ws.cell(2 + k, 1, f"={_q(DATA)}!$A${r}")
        ws.cell(2 + k, 3, f"={_q(DATA)}!$B${r}")
        ws.cell(2 + k, 2).fill = _INPUT
    for c, w in zip("ABC", (12, 60, 28)):
        ws.column_dimensions[c].width = w
    ws.freeze_panes = "A2"


def _check_sheet(wb, students, exc, last: int):
    """要確認の件数（B2）と一覧。一覧は右の作業列（隠す）で1件ずつ判定し、見つかったものだけを上から詰めて並べる。"""
    ws = wb.create_sheet(CHECK)
    ws["A1"] = "確認（「要確認」が0件になるまで通知表を印刷しない）"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"], ws["A3"] = "要確認の件数", "うち通知表データの要確認の欄"
    ws["A2"].font = ws["B2"].font = Font(bold=True, size=14, color="C00000")

    # 作業列: H 学籍番号, I 科目, J 何がおかしいか, K 通し番号（問題がある行だけ）
    work = []
    for i, (dept, sheet, src, sid, prefix, by_t) in enumerate(students):
        r = FIRST + i
        sid_ref = f"{_q(DATA)}!$A${r}"
        for t in sorted(by_t):
            name = TEMPLATE[t - 1][0]

            def cell(f, t=t, r=r):
                return f"{_q(CALC)}!${L(_calc_col(f, t))}${r}"

            msg = (f'=IF({cell("受講")}<>1,"",TRIM(IF({cell("点数")}="{UNSURE}","点数が数値でない ","")'
                   f'&IF(AND({cell("例外")}<>1,{cell("評定")}="{UNSURE}"),"評定がA〜Eでない ","")'
                   f'&IF({cell("授業数")}=0,"出席簿で授業数を数えられない（0回）","")))')
            work.append((f"={sid_ref}", name, msg))
        taken = _calc_row_range("受講", r)
        work.append((f"={sid_ref}", "（全科目）", f'=IF(SUM({taken})=0,"受けている科目が1つも無い","")'))
    for r in range(2, 2 + EXCEPTION_ROWS):
        ref = f"{_q(EXCEPTION)}!"
        work.append((f'={ref}$A{r}&""', f'={ref}$B{r}&""',
                     f'=IF(LEFT({ref}$G{r},3)="{UNSURE}","例外シート{r}行目："&MID({ref}$G{r},5,100),"")'))
    top, bottom = 6, 6 + len(work) - 1
    for c, h in zip("HIJK", ("作業: 学籍番号", "科目", "内容", "通し番号")):
        ws[f"{c}5"] = h
    for k, (sid, name, msg) in enumerate(work):
        r = top + k
        ws[f"H{r}"], ws[f"I{r}"], ws[f"J{r}"] = sid, name, msg
        ws[f"K{r}"] = f'=IF(J{r}<>"",{k + 1},"")'
    keys = f"$K${top}:$K${bottom}"
    ws["B2"] = f"=COUNT({keys})"
    ws["B3"] = f'=SUM({_q(DATA)}!${L(tail_column("要確認の件数"))}${FIRST}:${L(tail_column("要確認の件数"))}${last})'
    for c, h in zip("ABCD", ("番号", "学籍番号", "科目", "何がおかしいか")):
        ws[f"{c}5"] = h
    _head([ws[f"{c}5"] for c in "ABCD"])
    for k in range(1, len(work) + 1):
        r = top + k - 1
        pick = f"SMALL({keys},{k})"
        ws[f"A{r}"] = f'=IF({k}<=$B$2,{k},"")'
        for c, src in zip("BCD", "HIJ"):
            ws[f"{c}{r}"] = f'=IF({k}<=$B$2,INDEX(${src}${top}:${src}${bottom},{pick}),"")'
    for c in "HIJK":
        ws.column_dimensions[c].hidden = True
    for c, w in zip("ABCD", (8, 14, 30, 60)):
        ws.column_dimensions[c].width = w
    ws.freeze_panes = "A6"


# ---------- 照合（ADR 0002: 渡す前に一度、式の結果を report_card.py の計算と全員分比べる） ----------

def expected_cards(cards, exceptions: Sequence[tuple[str, str, str, str]] = ()):
    """read_cards の結果に例外処置を当てたもの（例外の行は withheld にし、評定欄に指定の文字）。"""
    import dataclasses
    marks = {(sid.strip(), subject.strip()): grade for sid, subject, grade, _ in exceptions}
    out = []
    for c in cards:
        lines = tuple(dataclasses.replace(l, withheld=True, withheld_grade=marks[(c.student_id, l.name)])
                      if (c.student_id, l.name) in marks else l for l in c.lines)
        out.append(dataclasses.replace(c, lines=lines))
    return out


def compare(ws, cards, tolerance: float = 1e-9) -> list[tuple[str, str, str]]:
    """ws: 再計算して値で読んだ「通知表データ」。cards: expected_cards の結果。
    食い違いを（学籍番号, 科目, 項目）で返す（値そのものは返さない。個人情報を出さないため）。"""
    def blank(v):
        return None if v == "" else v

    def near(a, b):
        return isinstance(a, (int, float)) and not isinstance(a, bool) and abs(a - float(b)) <= tolerance

    out = []
    last = FIRST + len(cards) - 1
    extra = [r for r in range(last + 1, ws.max_row + 1) if blank(ws.cell(r, 1).value) is not None]
    if extra:
        out.append(("", "", f"余分な行 {len(extra)}"))
    for k, c in enumerate(cards):
        r = FIRST + k
        sid = c.student_id
        if ws.cell(r, 1).value != sid:
            out.append((sid, "", "学籍番号（行の順）"))
            continue
        if (ws.cell(r, 2).value, ws.cell(r, 3).value, ws.cell(r, 4).value) != (c.name_en, c.name_ja, c.dept):
            out.append((sid, "", "氏名・カタカナ・学科"))
        for j in range(1, SLOTS + 1):
            got = {f: blank(ws.cell(r, slot_column(j, f)).value) for f in SLOT_FIELDS}
            if j > len(c.lines):
                if any(v is not None for v in got.values()):
                    out.append((sid, f"{j}枠目", "空のはずの枠に値"))
                continue
            line = c.lines[j - 1]
            if got["科目"] != line.name:
                out.append((sid, f"{j}枠目", "科目（順）"))
                continue
            if got["形態"] != line.kind:
                out.append((sid, line.name, "形態"))
            if got["設定"] != line.credits:
                out.append((sid, line.name, "設定"))
            if line.withheld:
                if got["点数"] is not None:
                    out.append((sid, line.name, "点数（例外は空欄）"))
                if (got["評定"] or "") != line.withheld_grade:
                    out.append((sid, line.name, "評定（例外）"))
            else:
                if not near(got["点数"], line.score):
                    out.append((sid, line.name, "点数"))
                if got["評定"] != line.grade:
                    out.append((sid, line.name, "評定"))
            if got["取得"] != line.earned:
                out.append((sid, line.name, "取得"))
        tail = {n: ws.cell(r, tail_column(n)).value for n in TAIL}
        if tail["前期設定計"] != c.credits_set:
            out.append((sid, "", "前期設定計"))
        if tail["前期取得計"] != c.credits_earned:
            out.append((sid, "", "前期取得計"))
        if tail["授業時数"] != c.hours:
            out.append((sid, "", "授業時数"))
        if not near(tail["出席時数"], c.attended):
            out.append((sid, "", "出席時数"))
        if not near(tail["出席率"], c.rate):
            out.append((sid, "", "出席率"))
        if tail["不合格"] is not any(l.grade == "E" or l.withheld for l in c.lines):
            out.append((sid, "", "不合格"))
        if tail["要確認の件数"] != 0:
            out.append((sid, "", "要確認の件数"))
    return out
