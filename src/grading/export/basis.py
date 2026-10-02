"""先生ごとの計算根拠（1枚の表）。学生を行、科目ごとに 授業数・欠席・遅刻・出席率・出席点・態度点 を並べる。

出席率は丸めずに分数で書く（例 17/18）。遅刻3回で欠席1回として 1−（欠席＋遅刻÷3）÷授業数。
"""

from collections.abc import Sequence
from dataclasses import dataclass
from fractions import Fraction

from openpyxl.styles import Alignment, Font, PatternFill

HEAD = ["授業数", "欠席", "遅刻", "出席率", "出席点", "態度点"]
_FILL = PatternFill("solid", fgColor="DDE7F3")


@dataclass(frozen=True)
class BasisRow:
    student_id: str
    name: str
    subject: str
    sessions: int
    absent: int
    late: int
    rate: Fraction
    attendance: Fraction
    attitude: Fraction | None     # 計算しない科目は None


def _num(v: Fraction | None):
    if v is None:
        return None
    return int(v) if v.denominator == 1 else float(v)


def add_teacher_basis(wb, teacher: str, subjects: Sequence[str], rows: Sequence[BasisRow]):
    head = HEAD if any(r.attitude is not None for r in rows) else HEAD[:-1]
    ws = wb.create_sheet(f"計算根拠_{teacher}")
    ws.cell(1, 1, "学籍番号")
    ws.cell(1, 2, "氏名")
    for i, subject in enumerate(subjects):
        col = 3 + i * len(head)
        ws.cell(1, col, subject)
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + len(head) - 1)
        for k, h in enumerate(head):
            ws.cell(2, col + k, h)
    for c in range(1, 3 + len(subjects) * len(head)):
        for r in (1, 2):
            cell = ws.cell(r, c)
            cell.font, cell.fill, cell.alignment = Font(bold=True), _FILL, Alignment(horizontal="center")
    order = list(dict.fromkeys(r.student_id for r in rows))
    by = {(r.student_id, r.subject): r for r in rows}
    for n, sid in enumerate(order):
        line = 3 + n
        first = next(r for r in rows if r.student_id == sid)
        ws.cell(line, 1, sid)
        ws.cell(line, 2, first.name)
        for i, subject in enumerate(subjects):
            r = by.get((sid, subject))
            if r is None:
                continue
            col = 3 + i * len(head)
            rate = "1" if r.rate == 1 else f"{r.rate.numerator}/{r.rate.denominator}"
            for k, v in enumerate([r.sessions, r.absent, r.late, rate, _num(r.attendance), _num(r.attitude)][:len(head)]):
                ws.cell(line, col + k, v)
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 24
    ws.freeze_panes = "C3"
    return ws
