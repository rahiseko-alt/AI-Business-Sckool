"""前期の追試の、科目ごとの出席簿（Excel）。対象は通知表で評定 E の科目（例外処置で F などにした科目は入れない）。

1科目1シート（A4縦）。並びは学科→学籍番号。出欠・追試の点数・備考は手書き用の空欄。
先頭の「一覧」シートに科目ごとの人数を出す。
"""

from collections.abc import Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from grading.export.report_card import Card, _one_decimal

FONT = "游ゴシック"
THIN = Side(style="thin", color="FF000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = PatternFill("solid", fgColor="FFE8E8E8")
COLUMNS = (("No.", 5), ("学科", 14), ("学籍番号", 12), ("氏名", 34), ("カタカナ", 20),
           ("前期の点数", 9), ("出欠", 7), ("追試の点数", 9), ("備考", 14))


def targets(cards: Sequence[Card]) -> dict[str, list[tuple[Card, object]]]:
    """科目名 → [(学生, 前期の点数)]。評定 E の科目だけ。科目は通知表の順（最初に出た順）、学生は学科→学籍番号。"""
    out: dict[str, list] = {}
    for c in cards:
        for l in c.lines:
            out.setdefault(l.name, [])
            if not l.withheld and l.grade == "E":
                out[l.name].append((c, l.score))
    for rows in out.values():
        rows.sort(key=lambda x: (x[0].dept, x[0].student_id))
    return {k: v for k, v in out.items() if v}


def _sheet_title(name: str) -> str:
    for ch in "[]:*?/\\":
        name = name.replace(ch, "・")
    return name[:31]


def build(cards: Sequence[Card], year: str = "2026年度"):
    wb = Workbook()
    summary = wb.active
    summary.title = "一覧"
    t = targets(cards)
    summary["A1"] = f"{year} 前期 追試 対象者数（評定 E の科目）"
    summary["A1"].font = Font(name=FONT, bold=True, size=13)
    for c, h in enumerate(("科目", "人数"), 1):
        x = summary.cell(3, c, h)
        x.font, x.fill, x.border = Font(name=FONT, bold=True), HEAD, BOX
    for i, (subject, rows) in enumerate(t.items(), 4):
        summary.cell(i, 1, subject).border = BOX
        summary.cell(i, 2, len(rows)).border = BOX
    total = 4 + len(t)
    summary.cell(total, 1, "合計（のべ）").font = Font(name=FONT, bold=True)
    summary.cell(total, 2, sum(len(v) for v in t.values())).font = Font(name=FONT, bold=True)
    summary.column_dimensions["A"].width, summary.column_dimensions["B"].width = 34, 8

    for subject, rows in t.items():
        ws = wb.create_sheet(_sheet_title(subject))
        ws["A1"] = f"{year} 前期 追試 出席簿"
        ws["A1"].font = Font(name=FONT, bold=True, size=14)
        ws["A2"] = f"科目: {subject}"
        ws["A2"].font = Font(name=FONT, bold=True, size=12)
        ws["F2"] = "実施日:　　　年　　月　　日"
        ws["A3"] = f"対象 {len(rows)}人"
        ws["F3"] = "担当:"
        for cell in ("F2", "A3", "F3"):
            ws[cell].font = Font(name=FONT, size=10)
        for c, (h, w) in enumerate(COLUMNS, 1):
            x = ws.cell(5, c, h)
            x.font, x.fill, x.border = Font(name=FONT, bold=True, size=10), HEAD, BOX
            x.alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions[x.column_letter].width = w
        for i, (card, score) in enumerate(rows, 1):
            r = 5 + i
            values = (i, card.dept, card.student_id, card.name_en, card.name_ja,
                      "" if score is None else float(_one_decimal(score)), "", "", "")
            for c, v in enumerate(values, 1):
                x = ws.cell(r, c, v)
                x.border, x.font = BOX, Font(name=FONT, size=10)
                x.alignment = Alignment(vertical="center", shrink_to_fit=c in (4, 5),
                                        horizontal="center" if c in (1, 7) else None)
            ws.row_dimensions[r].height = 24
        ws.print_title_rows = "5:5"
        ws.page_setup.paperSize, ws.page_setup.orientation = ws.PAPERSIZE_A4, "portrait"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.print_options.horizontalCentered = True
        ws.oddFooter.center.text = "&P / &N"
    return wb
