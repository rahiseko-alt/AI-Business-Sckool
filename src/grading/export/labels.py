"""全員分の名前シール（A4・12面＝2列×6段）を Excel で作る。AI を使わずに直せるよう、中身は「名簿」シートの式で読む。

シール1枚の中身（3行）: 「2026年度　学科」／学籍番号／「英字の氏名　呼び名」。
呼び名は名簿の黄色の欄に入れる（空なら英字の氏名だけ）。名簿を直せばシールも変わる。
シールの面は作った人数を12面単位に切り上げた数だけ用意する（空いた面は何も出ない）。

寸法は A4 12面の定番（エーワン 72212 など）: 1面 86.4×42.3mm、上の余白 21.5mm、左の余白 18.6mm、すきま無し。
拡大縮小は100%固定（ずれると台紙の切れ目と合わなくなる）。
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.pagebreak import Break

ROSTER = "名簿"
LABELS = "シール"
YEAR = "2026年度"
MAX_STUDENTS = 120
COLS, ROWS = 2, 6                      # 1枚 12面
LABEL_W_MM, LABEL_H_MM = 86.4, 42.3
TOP_MM, LEFT_MM = 21.5, 18.6
# 1面を5行で作る（上の余白・3行・下の余白）。和は 42.3mm＝119.9pt
LINE_PT = (22, 25, 25, 28, 19.9)
FONT = "游ゴシック"
YELLOW = PatternFill("solid", fgColor="FFFFFF00")
HEAD = PatternFill("solid", fgColor="FFE8E8E8")
THIN = Side(style="thin", color="FF000000")


def _width(mm: float) -> float:
    """mm を列幅（既定のフォント Calibri 11、数字1文字 7px）に直す。"""
    return round((mm / 25.4 * 96 - 5) / 7, 2)


def build(students: list[tuple[str, str, str]]):
    """students: (学科, 学籍番号, 英字の氏名) の並び。シールはこの順に左上から右へ、上から下へ。"""
    if len(students) > MAX_STUDENTS:
        raise ValueError(f"{len(students)}人は多すぎる（最大 {MAX_STUDENTS}人）")
    wb = Workbook()
    labels = wb.active
    labels.title = LABELS
    roster = wb.create_sheet(ROSTER)

    # 名簿
    for c, (h, w) in enumerate((("学科", 18), ("学籍番号", 14), ("英字の氏名", 40), ("呼び名（ここに入れる）", 24)), 1):
        cell = roster.cell(1, c, h)
        cell.font, cell.fill = Font(name=FONT, bold=True), HEAD
        roster.column_dimensions["ABCD"[c - 1]].width = w
    for i, (dept, sid, name) in enumerate(students, 2):
        for c, v in enumerate((dept, sid, name), 1):
            roster.cell(i, c, v).font = Font(name=FONT)
    for i in range(2, MAX_STUDENTS + 2):
        roster.cell(i, 4).fill = YELLOW
        roster.cell(i, 4).font = Font(name=FONT)
    roster.freeze_panes = "A2"

    # シール（1面＝5行×1列。左の面が A 列、右の面が B 列）
    labels.column_dimensions["A"].width = labels.column_dimensions["B"].width = _width(LABEL_W_MM)
    labels.sheet_view.showGridLines = False
    per_page = COLS * ROWS
    pages = max(1, -(-len(students) // per_page))
    for k in range(pages * per_page):
        page, slot = divmod(k, per_page)
        r0 = 1 + (page * ROWS + slot // COLS) * len(LINE_PT)
        col = 1 + slot % COLS
        n = k + 2                                  # 名簿の行
        ref = lambda c: f"{ROSTER}!${c}${n}"
        lines = (f'=IF({ref("B")}="","","{YEAR}　"&{ref("A")})',
                 f'=IF({ref("B")}="","",{ref("B")})',
                 f'=IF({ref("B")}="","",{ref("C")}&IF({ref("D")}="","","　"&{ref("D")}))')
        for i, (text, size) in enumerate(zip(lines, (11, 14, 13)), 1):
            cell = labels.cell(r0 + i, col, text)
            cell.font = Font(name=FONT, size=size, bold=i > 1)
            cell.alignment = Alignment(horizontal="center", vertical="center", shrink_to_fit=True)
    for p in range(pages):
        for i, h in enumerate(LINE_PT * ROWS):
            labels.row_dimensions[1 + p * ROWS * len(LINE_PT) + i].height = h
        if p < pages - 1:
            labels.row_breaks.append(Break(id=(p + 1) * ROWS * len(LINE_PT)))
    labels.print_area = f"A1:B{pages * ROWS * len(LINE_PT)}"
    ps = labels.page_setup
    ps.paperSize, ps.orientation, ps.scale = labels.PAPERSIZE_A4, "portrait", 100
    m = labels.page_margins
    m.top, m.left = TOP_MM / 25.4, LEFT_MM / 25.4
    m.bottom = m.right = 0.1
    m.header = m.footer = 0
    labels.print_options.horizontalCentered = labels.print_options.verticalCentered = False
    return wb
