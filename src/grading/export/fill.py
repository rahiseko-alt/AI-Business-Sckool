"""白紙にした成績表（AI計算版の枠）に、システムが計算した値だけを入れる。

科目名（3行目）と評価項目名（4行目）と学籍番号（B列）からセルを探す。見つからなければ推測せずに止める。
入れた値はすべて「根拠」シートに、計算式と元のセルを添えて並べる。値は丸めない（割り切れない値は分数で残す）。
"""

from collections.abc import Sequence
from dataclasses import dataclass
from fractions import Fraction
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from grading.export.provenance import format_number
from grading.importing.attendance import unify_subject

SUBJECT_ALIASES = {"AI演習(実践)": "AI演習", "国際理解": "国際社会I:異文化理解"}   # 利用者の回答（質問13）で確定


@dataclass(frozen=True)
class FilledValue:
    student_id: str
    subject: str
    item: str
    value: Fraction
    explanation: str
    evidence: tuple[str, ...]


def _subject(name: str) -> str:
    n = unify_subject(name)
    return SUBJECT_ALIASES.get(n, n)


def _columns(ws) -> dict[tuple[str, str], int]:
    out, current = {}, None
    for c in range(5, ws.max_column + 1):
        head = ws.cell(3, c).value
        if isinstance(head, str) and head.strip():
            current = _subject(head)
        item = ws.cell(4, c).value
        if current and isinstance(item, str) and item.strip():
            out[(current, unify_subject(item))] = c
    return out


def fill_template(src: str | Path, dst: str | Path, values: Sequence[FilledValue]) -> Path:
    wb = load_workbook(src)
    ws = wb.worksheets[0]
    columns = _columns(ws)
    rows = {str(ws.cell(r, 2).value).strip(): r for r in range(6, ws.max_row + 1) if ws.cell(r, 2).value}
    if "根拠" in wb.sheetnames:
        ev = wb["根拠"]
    else:
        ev = wb.create_sheet("根拠")
        ev.append(["学籍番号", "科目", "評価項目", "セル", "値", "正確な値", "計算", "元のセル"])
        for cell in ev[1]:
            cell.font = Font(bold=True)
    written = set()
    for v in values:
        key = (_subject(v.subject), unify_subject(v.item))
        if key not in columns:
            raise KeyError(f"成績表に「{v.subject}」の「{v.item}」の列が無い")
        if v.student_id not in rows:
            raise KeyError(f"成績表に学籍番号 {v.student_id} が無い")
        r, c = rows[v.student_id], columns[key]
        if (r, c) in written:
            raise ValueError(f"{v.student_id} {v.subject} {v.item} に2回書こうとした")
        written.add((r, c))
        number = int(v.value) if v.value.denominator == 1 else float(v.value)
        ws.cell(r, c).value = number
        ev.append([v.student_id, v.subject, v.item, f"{get_column_letter(c)}{r}", number,
                   format_number(v.value).split("（")[0], v.explanation, "、".join(v.evidence)])
    for col, width in zip("ABCDEFGH", (12, 22, 12, 7, 8, 10, 40, 60)):
        ev.column_dimensions[col].width = width
    ev.freeze_panes = "A2"
    dst = Path(dst)
    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)
    return dst


AI_FILL = PatternFill("solid", fgColor="DDEBF7")
EVIDENCE_HEADER = ["学籍番号", "科目", "評価項目", "セル", "原本の値", "AIの値", "正確な値", "計算", "元のセル"]


def build_ai_workbook(original: str | Path, dst: str | Path, values: Sequence[FilledValue],
                      teacher_of: dict[str, str]) -> Path:
    """原本を丸ごと写し、AIが計算したセルだけを置き換えた「AI計算版」を作る。

    シート: AI計算版（AIのセルは色付き、ほかは原本のまま。合計などの式も残る）／原本（手を付けない）／
    計算根拠_先生名（先生ごとに、原本の値・AIの値・計算・元のセル）。
    """
    wb = load_workbook(original)
    cached = load_workbook(original, data_only=True).worksheets[0]
    ai = wb.worksheets[0]
    keep = wb.copy_worksheet(ai)
    keep.title = "原本"
    ai.title = "AI計算版"
    if isinstance(ai["A1"].value, str):
        ai["A1"].value += "（AI計算版）"
    columns = _columns(ai)
    rows = {str(ai.cell(r, 2).value).strip(): r for r in range(6, ai.max_row + 1) if ai.cell(r, 2).value}
    sheets: dict[str, object] = {}
    written = set()
    for v in values:
        key = (_subject(v.subject), unify_subject(v.item))
        if key not in columns:
            raise KeyError(f"成績表に「{v.subject}」の「{v.item}」の列が無い")
        if v.student_id not in rows:
            raise KeyError(f"成績表に学籍番号 {v.student_id} が無い")
        teacher = teacher_of[_subject(v.subject)]
        r, c = rows[v.student_id], columns[key]
        if (r, c) in written:
            raise ValueError(f"{v.student_id} {v.subject} {v.item} に2回書こうとした")
        written.add((r, c))
        before = cached.cell(r, c).value
        if before is None:
            before = ai.cell(r, c).value
        number = int(v.value) if v.value.denominator == 1 else float(v.value)
        ai.cell(r, c).value = number
        ai.cell(r, c).fill = AI_FILL
        if teacher not in sheets:
            ev = wb.create_sheet(f"計算根拠_{teacher}")
            ev.append(EVIDENCE_HEADER)
            for cell in ev[1]:
                cell.font = Font(bold=True)
            for col, width in zip("ABCDEFGHI", (12, 22, 10, 7, 10, 8, 10, 60, 60)):
                ev.column_dimensions[col].width = width
            ev.freeze_panes = "A2"
            sheets[teacher] = ev
        sheets[teacher].append([v.student_id, v.subject, v.item, f"{get_column_letter(c)}{r}", before, number,
                                format_number(v.value).split("（")[0], v.explanation, "、".join(v.evidence)])
    dst = Path(dst)
    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)
    return dst
