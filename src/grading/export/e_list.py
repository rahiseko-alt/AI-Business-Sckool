"""科目ごとの評定E一覧。再計算済みの成績表（値）から読む。

その学生の評価項目がすべて空欄の科目（受講していない科目）は一覧に入れない。
"""

from collections.abc import Sequence
from dataclasses import dataclass

from openpyxl.styles import Font, PatternFill


@dataclass(frozen=True)
class GradeRow:
    subject: str
    student_id: str
    name: str
    total: object
    grade: object


def grade_rows(ws) -> list[GradeRow]:
    groups, current = [], None
    for c in range(5, ws.max_column + 1):
        head = ws.cell(3, c).value
        if isinstance(head, str) and head.strip():
            current = {"subject": head.strip(), "items": [], "total": None, "grade": None}
            groups.append(current)
        item = ws.cell(4, c).value
        if current is None or not isinstance(item, str):
            continue
        if item.strip() == "合計":
            current["total"] = c
        elif item.strip() == "評定":
            current["grade"] = c
        elif current["total"] is None:
            current["items"].append(c)
    out = []
    for g in groups:
        if g["grade"] is None:
            continue
        for r in range(6, ws.max_row + 1):
            sid = ws.cell(r, 2).value
            if not sid:
                continue
            if all(ws.cell(r, c).value in (None, "") for c in g["items"]):
                continue
            out.append(GradeRow(g["subject"], str(sid).strip(), ws.cell(r, 3).value,
                                ws.cell(r, g["total"]).value if g["total"] else None, ws.cell(r, g["grade"]).value))
    return out


def add_e_list(wb, rows: Sequence[GradeRow], title: str = "E一覧"):
    ws = wb.create_sheet(title)
    ws.append(["科目", "学籍番号", "氏名", "合計", "評定"])
    for cell in ws[1]:
        cell.font, cell.fill = Font(bold=True), PatternFill("solid", fgColor="DDE7F3")
    for r in rows:
        if r.grade == "E":
            ws.append([r.subject, r.student_id, r.name, r.total, r.grade])
    for col, width in zip("ABCDE", (28, 12, 30, 8, 6)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    return ws
