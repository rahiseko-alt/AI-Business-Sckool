"""全員分の名前シール（A4・12面）の Excel を作る。

    python3 scripts/build_name_labels.py [成績表.xlsx]

入力を省くと data/input/前期成績表.xlsx の AI計算版シート（国際→総合の順）から、学科・学籍番号・英字の氏名を読む。
出力は data/output/名前シール_2026.xlsx（個人情報のため Git には入れない）。呼び名は出力の「名簿」シートの黄色の欄に入れる。
"""

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook  # noqa: E402

from grading.export.labels import build  # noqa: E402

INPUT = ROOT / "data" / "input" / "前期成績表.xlsx"
OUTPUT = ROOT / "data" / "output" / "名前シール_2026.xlsx"
DEPTS = (("AI計算版_国際", "国際ビジネス科"), ("AI計算版_総合", "総合ビジネス科"))


def students(source: Path):
    wb = load_workbook(source, read_only=True)
    out = []
    for sheet, dept in DEPTS:
        for row in wb[sheet].iter_rows(min_row=6, max_col=3, values_only=True):
            if isinstance(row[1], str) and row[1].startswith("AIBC"):
                out.append((dept, row[1].strip(), " ".join(str(row[2]).split())))
    return out


if __name__ == "__main__":
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else INPUT
    s = students(src)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    build(s).save(OUTPUT)
    print(f"{len(s)}人分 → {OUTPUT}")
