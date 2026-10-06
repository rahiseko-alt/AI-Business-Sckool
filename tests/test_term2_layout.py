import shutil
import subprocess

import pytest
from openpyxl import Workbook, load_workbook

from grading.term2.layout import (FIRST_ROW, LINKS_SHEET, SUBMISSION_TAB, SUBMITTED_CELL, gradebook_cells,
                                  submission_cells)


def _sub():
    return submission_cells("試作科目", "担当先生", ["TEST01", "TEST02"])


def test_submission_has_points_submitted_and_student_rows():
    cells = _sub()
    assert cells.values[0][:2] == ["科目名", "試作科目"] and cells.values[0][3] == "担当先生"
    assert cells.values[2][0] == "学籍番号" and cells.values[2][-1] == "合計"
    assert [r[0] for r in cells.values[FIRST_ROW - 1:]] == ["TEST01", "TEST02"]


def test_yellow_is_only_points_submitted_and_manual_columns():
    cells = _sub()
    assert set(cells.yellow) == {"E2", "H2", "K2", SUBMITTED_CELL, "D4:D5", "G4:G5", "J4:J5"}
    assert all(cells.values[1][ord(c) - ord("A")] is None for c in "EHK")   # 配点の黄の欄は空で始まる


def test_red_rules_cover_points_total_and_every_adopted_column():
    rules = dict(_sub().red)
    assert "L2" in rules and "100" in rules["L2"]
    for col in "EHK":
        rng = f"{col}4:{col}5"
        assert rng in rules
        f = rules[rng]
        assert "ISNUMBER" in f and f"{col}$2" in f and SUBMITTED_CELL.replace("F", "$F$") in f


def test_gradebook_reads_adopted_values_and_submitted_through_links_only():
    cells = gradebook_cells(["試作科目"], rows=2)
    url_cell = f"{LINKS_SHEET}!$B${FIRST_ROW}"
    formulas = [f for row in cells.tabs["集計"] for f in row if isinstance(f, str) and f.startswith("=")]
    assert len(formulas) == 2 and all(url_cell in f for f in formulas)
    assert all("docs.google.com" not in f for f in formulas)       # ファイルを直接名指ししない
    assert any(f"'{SUBMISSION_TAB}'!{SUBMITTED_CELL}" in f for f in formulas)
    assert any("1,5,8,11,12" in f for f in formulas)                # 学籍番号・採用値3つ・合計
    assert cells.tabs[LINKS_SHEET][3][0] == "試作科目" and cells.yellow[LINKS_SHEET] == ["B4"]


def _has_calc() -> bool:
    return shutil.which("soffice") is not None and subprocess.run(
        ["dpkg", "-s", "libreoffice-calc"], capture_output=True).returncode == 0


@pytest.mark.skipif(not _has_calc(), reason="LibreOffice Calc が無いため式の再計算を確認できない")
def test_manual_input_wins_over_auto_and_falls_back_when_cleared(tmp_path):
    cells = _sub()
    wb = Workbook()
    ws = wb.active
    for r, row in enumerate(cells.values, 1):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    ws["C4"], ws["D4"] = 30, 25        # 自動30・手入力25 → 25
    ws["C5"], ws["D5"] = 30, None      # 手入力なし → 自動の30
    ws["G4"], ws["J4"] = 20, 40
    src = tmp_path / "s.xlsx"
    wb.save(src)
    out = tmp_path / "out"
    subprocess.run(["soffice", f"-env:UserInstallation=file://{tmp_path}/profile", "--headless", "--convert-to",
                    "xlsx:Calc MS Excel 2007 XML", "--outdir", str(out), str(src)],
                   check=True, capture_output=True, timeout=180)
    got = load_workbook(out / "s.xlsx", data_only=True).active
    assert got["E4"].value == 25 and got["E5"].value == 30
    assert got["L4"].value == 85
