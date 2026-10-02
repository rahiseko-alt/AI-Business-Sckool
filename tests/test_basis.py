from fractions import Fraction

from openpyxl import Workbook, load_workbook

from grading.export.basis import BasisRow, add_teacher_basis


def test_teacher_basis_is_one_compact_table(tmp_path):
    wb = Workbook()
    rows = [BasisRow("AIBC26001", "TARO", "マーケティング", 30, 1, 2, Fraction(1) - Fraction(5, 90), Fraction(9), Fraction(6)),
            BasisRow("AIBC26001", "TARO", "AI演習", 30, 0, 0, Fraction(1), Fraction(10), None)]
    add_teacher_basis(wb, "百井", ["マーケティング", "AI演習"], rows)
    wb.save(tmp_path / "b.xlsx")
    ws = load_workbook(tmp_path / "b.xlsx")["計算根拠_百井"]
    assert [ws.cell(1, c).value for c in (1, 3, 9)] == ["学籍番号", "マーケティング", "AI演習"]
    assert [ws.cell(2, c).value for c in range(3, 9)] == ["授業数", "欠席", "遅刻", "出席率", "出席点", "態度点"]
    assert [ws.cell(3, c).value for c in range(1, 9)] == ["AIBC26001", "TARO", 30, 1, 2, "17/18", 9, 6]
    assert ws.cell(3, 14).value is None and ws.cell(3, 13).value == 10


def test_attitude_column_is_left_out_when_not_calculated(tmp_path):
    wb = Workbook()
    add_teacher_basis(wb, "樋口", ["ビジネス日本語"],
                      [BasisRow("AIBC26001", "TARO", "ビジネス日本語", 30, 0, 0, Fraction(1), Fraction(10), None)])
    ws = wb["計算根拠_樋口"]
    assert [ws.cell(2, c).value for c in range(3, 9)] == ["授業数", "欠席", "遅刻", "出席率", "出席点", None]
