"""通知表データのシート（式）を小さな成績表で作り、LibreOffice で再計算して report_card.py の計算と比べる。

LibreOffice が無いと再計算の確認ができないので、飛ばさずに失敗させる（ADR 0002: 渡す前に式の結果を照合する）。
"""

import datetime as dt
import shutil
import subprocess
from fractions import Fraction

import pytest
from openpyxl import Workbook, load_workbook

from grading.export.report_card import read_cards
from grading.export.report_card_data import (CHECK, DATA, EXCEPTION, REMARK, add_report_card_data, compare,
                                             expected_cards, slot_column, tail_column)
from grading.importing.attendance import read_register, subject_counts

DEPTS = (("AI計算版_国際", "国際ビジネス科", "国際ビジネスAI科"), ("AI計算版_総合", "総合ビジネス科", "総合ビジネス科"))
LOOKUP = '=LOOKUP({0},{{0,60,70,80,90}},{{"E","D","C","B","A"}})'


def _ai(wb, title, students, subjects):
    """AI計算版と同じ並び。subjects: [(科目名, 評価項目の数)]。students: [(学籍番号, 英字, カナ, {科目: [点]})]。"""
    ws = wb.create_sheet(title)
    ws["B3"] = "学籍番号"
    c, layout = 5, {}
    for name, n in subjects:
        ws.cell(3, c, name)
        for k in range(n):
            ws.cell(4, c + k, f"項目{k + 1}")
        ws.cell(4, c + n, "合計")
        ws.cell(4, c + n + 1, "評定")
        layout[name] = (c, n)
        c += n + 2
    for i, (sid, en, kana, scores) in enumerate(students):
        r = 6 + i
        ws.cell(r, 2, sid), ws.cell(r, 3, en), ws.cell(r, 4, kana)
        for name, (c, n) in layout.items():
            for k, v in enumerate(scores.get(name, [])):
                ws.cell(r, c + k, v)
            first, last = ws.cell(r, c).coordinate, ws.cell(r, c + n - 1).coordinate
            total = ws.cell(r, c + n).coordinate
            ws.cell(r, c + n, f"=SUM({first}:{last})")
            ws.cell(r, c + n + 1, LOOKUP.format(total))
    return ws


def _register(wb, title, students, heads, marks):
    """出席簿の月別シート。heads: [(3行目, 4行目)]（E列から）。marks: {(学籍番号, 列の番号0始まり): 記号}。"""
    ws = wb.create_sheet(title)
    for k, (date, subject) in enumerate(heads):
        ws.cell(3, 5 + k, date)
        ws.cell(4, 5 + k, subject)
        ws.cell(5, 5 + k, k + 1)
    ws.cell(3, 5 + len(heads), "授業合計数")
    for i, sid in enumerate(students):
        ws.cell(6 + i, 2, sid)
        for k in range(len(heads)):
            ws.cell(6 + i, 5 + k, marks.get((sid, k)))
    return ws


def _book():
    """国際2人・総合1人。国際の5月は日付の無い見出し（数えない列）、康熙部首の科目名、国際社会の別名を含む。"""
    wb = Workbook()
    wb.remove(wb.active)
    _ai(wb, "AI計算版_国際", [
        ("AIBC26001", "TARO  YAMADA", "タロウ　ヤマ\nダ", {"ビジネス日本語": [30, 25.25], "国際理解": [40, 15]}),
        ("AIBC26002", "HANA SATO", "ハナ", {"ビジネス日本語": [45, 46.5], "日本語能力強化演習": [50, 20], "国際理解": [30, 41]}),
    ], [("ビジネス日本語", 2), ("国際理解", 2), ("日本語能力強化演習", 2)])
    _ai(wb, "AI計算版_総合", [
        ("AIBC26003", "JIRO", "ジロウ", {"ビジネス日本語": [35, 36], "総合ビジネス概論": [None, 70]}),
    ], [("ビジネス日本語", 2), ("総合ビジネス概論", 2)])
    kokusai = ["AIBC26001", "AIBC26002"]
    _register(wb, "国際ビジネスAI科 4月", kokusai,
              [("4月 13日（月）", "ビジネス日本語"), (None, "ビジネス日本語"), ("4月 14日（火）", "国際社会Ⅰ：異⽂化理解"),
               (None, "⽇本語能⼒強化演習"), (None, "⽇本語能⼒強化演習 ")],
              {("AIBC26001", 0): "欠", ("AIBC26001", 1): "遅", ("AIBC26001", 2): "×", ("AIBC26002", 3): "遅",
               ("AIBC26002", 4): "欠", ("AIBC26001", 3): "×", ("AIBC26001", 4): "×"})
    _register(wb, "国際ビジネスAI科 5月", kokusai,
              [("　月　　日（月）", "ビジネス日本語"), (None, "国際社会Ⅰ：異⽂化理解"),
               (dt.datetime(2026, 5, 11), "ビジネス日本語"), ("=TEXT(G3,\"( aaa )\")", "国際社会Ⅰ：異⽂化理解"),
               ("=G3+1", "ビジネス日本語")],
              {("AIBC26001", 0): "欠", ("AIBC26002", 1): "欠", ("AIBC26002", 2): "遅", ("AIBC26001", 3): "遅",
               ("AIBC26002", 4): "欠"})
    _register(wb, "総合ビジネス科 4月", ["AIBC26003"],
              [("4月 13日（月）", "ビジネス日本語 "), (None, "総合ビジネス概論"), ("4月 14日（火）", "ビジネス日本語")],
              {("AIBC26003", 2): "欠"})
    return wb


def _recalc(path, tmp_path):
    """LibreOffice で再計算した値のブック。LibreOffice が無ければ失敗させる（飛ばさない）。"""
    assert shutil.which("soffice"), "LibreOffice（soffice）が無いため、式の結果を確認できない"
    out = tmp_path / "out"
    subprocess.run(["soffice", f"-env:UserInstallation=file://{tmp_path}/profile", "--headless", "--convert-to",
                    "xlsx:Calc MS Excel 2007 XML", "--outdir", str(out), str(path)],
                   check=True, capture_output=True, timeout=300)
    return load_workbook(out / path.name, data_only=True)


def _expected(source, values, exceptions=()):
    """report_card.py の計算（照合の正）。source: 成績表のファイル、values: その再計算した値。"""
    cards = []
    for sheet, dept, prefix in DEPTS:
        cards += read_cards(values[sheet], dept, subject_counts(read_register(source, sheet_prefix=prefix)))
    return expected_cards(cards, exceptions)


def _build(tmp_path, edit=None, exceptions=(), name="built.xlsx"):
    """編集（edit(wb)）を当てた成績表と、それに通知表データを足したものを保存し、両方の値と照合の正を返す。"""
    wb = _book()
    if edit:
        edit(wb)
    source = tmp_path / "source.xlsx"
    wb.save(source)
    add_report_card_data(wb, DEPTS, exceptions)
    built = tmp_path / name
    wb.save(built)
    values = _recalc(built, tmp_path)
    return values, _expected(source, values, exceptions)


def test_sheets_follow_the_contract_layout_and_hold_only_formulas():
    wb = _book()
    assert add_report_card_data(wb, DEPTS) == 3
    ws = wb[DATA]
    assert [ws.cell(3, c).value for c in range(1, 5)] == ["学籍番号", "氏名", "カタカナ", "学科"]
    assert (slot_column(1, "科目"), slot_column(13, "取得")) == (5, 82)
    assert ws["E3"].value == "科目1" and ws["CD3"].value == "取得13"
    assert [ws.cell(3, c).value for c in range(tail_column("前期設定計"), tail_column("特記事項") + 1)] == [
        "前期設定計", "前期取得計", "授業時数", "出席時数", "出席率", "不合格", "要確認の件数", "特記事項"]
    assert tail_column("前期設定計") == 83 and tail_column("特記事項") == 90       # CE〜CL
    for row in ws.iter_rows(min_row=4, max_row=6):
        for cell in row:
            assert cell.column == 4 or (isinstance(cell.value, str) and cell.value.startswith("=")), cell.coordinate
    assert [c.value for c in wb[EXCEPTION][1]][:4] == ["学籍番号", "科目（通知表の科目名）", "表示する評定", "理由"]
    assert wb[REMARK]["A1"].value == "学籍番号" and wb[CHECK]["B2"].value.startswith("=")
    with pytest.raises(ValueError, match="もう有る"):
        add_report_card_data(wb, DEPTS)


def test_recalculated_data_matches_read_cards_for_every_student(tmp_path):
    values, cards = _build(tmp_path, exceptions=(("AIBC26002", "日本語能力強化演習", "F", "利用者の指示による例外処置"),))
    assert compare(values[DATA], cards) == []
    ws = values[DATA]
    # 国際1人目: 受けていない日本語能力強化演習は載らず、2科目が上から詰まる
    assert [ws.cell(4, slot_column(j, "科目")).value for j in (1, 2)] == ["ビジネス日本語Ⅰ", "国際社会・異文化理解Ⅰ"]
    assert ws.cell(4, slot_column(3, "科目")).value in (None, "")
    assert ws.cell(4, 2).value == "TARO YAMADA" and ws.cell(4, 3).value == "タロウ ヤマ ダ"
    # 例外の行: 点数・取得は空欄、評定は F、不合格になる
    j = [ws.cell(5, slot_column(k, "科目")).value for k in range(1, 4)].index("日本語能力強化演習") + 1
    assert [ws.cell(5, slot_column(j, f)).value for f in ("点数", "評定", "設定")] in (["", "F", 3], [None, "F", 3])
    assert ws.cell(5, tail_column("不合格")).value is True
    assert ws.cell(4, tail_column("特記事項")).value == "該当なし"
    assert values[CHECK]["B2"].value == 0
    assert values[EXCEPTION]["G2"].value == "反映済み"


def test_editing_scores_register_and_exceptions_changes_the_data_like_read_cards(tmp_path):
    (tmp_path / "before").mkdir()
    (tmp_path / "after").mkdir()
    _, before = _build(tmp_path / "before")
    assert [l.grade for l in before[0].lines] == ["E", "E"]

    def edit(wb):
        wb["AI計算版_国際"]["E6"] = 35                       # ビジネス日本語 55.25 → 60.25（E → D）
        wb["AI計算版_国際"]["I6"] = 50                       # 国際理解 55 → 65（E → D）
        wb["国際ビジネスAI科 5月"]["G7"] = "欠"               # 2人目の遅を欠にする
        wb["総合ビジネス科 4月"]["E6"] = "欠"                 # 3人目に欠を1つ足す

    exceptions = (("AIBC26003", "ビジネス日本語Ⅰ", "F", "照合用"),)
    values, after = _build(tmp_path / "after", edit, exceptions)
    assert [l.grade for l in after[0].lines] == ["D", "D"]
    assert after[1].attended == before[1].attended - 1 + Fraction(1, 3)      # 遅（1/3）→ 欠（1）
    assert after[2].attended == before[2].attended - 1 and after[2].lines[0].withheld
    assert compare(values[DATA], after) == []
    ws = values[DATA]
    assert ws.cell(4, tail_column("不合格")).value is False and ws.cell(6, tail_column("不合格")).value is True
    assert ws.cell(4, slot_column(1, "取得")).value == 2


def test_undeterminable_values_show_unsure_and_are_listed(tmp_path):
    def edit(wb):
        wb["AI計算版_国際"]["H6"] = "可"                      # 評定が A〜E でない
        wb["総合ビジネス科 4月"]["E4"] = "別の科目"            # ビジネス日本語の列を2つとも別の科目にすると
        wb["総合ビジネス科 4月"]["G4"] = "別の科目"            # ビジネス日本語の授業数が0になる

    wb = _book()
    edit(wb)
    add_report_card_data(wb, DEPTS, (("AIBC26001", "存在しない科目", "F", ""), ("AIBC99999", "ビジネス日本語Ⅰ", "F", ""),
                                     ("AIBC26001", "日本語能力強化演習", "F", "")))
    wb[REMARK]["B3"], wb[REMARK]["B4"] = " 休学中 ", "　"     # 全角の空白だけなら書いていないのと同じ
    path = tmp_path / "unsure.xlsx"
    wb.save(path)
    values = _recalc(path, tmp_path)
    ws = values[DATA]
    assert ws.cell(4, slot_column(1, "評定")).value == "要確認"
    assert ws.cell(4, slot_column(1, "取得")).value == "要確認"
    assert ws.cell(4, tail_column("前期取得計")).value == "要確認"
    assert ws.cell(4, tail_column("不合格")).value is True           # 国際理解が E なので、評定が1つ分からなくても不合格
    assert ws.cell(4, tail_column("要確認の件数")).value == 3
    assert [ws.cell(6, tail_column(n)).value for n in ("授業時数", "出席時数", "出席率")] == ["要確認"] * 3
    assert [ws.cell(r, tail_column("特記事項")).value for r in (5, 6)] == ["休学中", "該当なし"]
    check = values[CHECK]
    listed = [(check.cell(r, 2).value, check.cell(r, 3).value, check.cell(r, 4).value) for r in range(6, 6 + check["B2"].value)]
    assert check["B2"].value == 5 and check["B3"].value == 6
    assert ("AIBC26001", "ビジネス日本語Ⅰ", "評定がA〜Eでない") in listed
    assert ("AIBC26003", "ビジネス日本語Ⅰ", "出席簿で授業数を数えられない（0回）") in listed
    assert sum("例外シート" in m for *_, m in listed) == 3
    assert check.cell(6 + check["B2"].value, 2).value in (None, "")
