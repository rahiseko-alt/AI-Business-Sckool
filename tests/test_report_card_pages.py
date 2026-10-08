"""通知表_全員・通知表_1人・成績証明書_1人（通知表データを式で読む印刷用のシート）。

通知表データを値で置いた小さなブックに add_pages し、LibreOffice で再計算して印刷面の値を確かめる。
「1人1ページ」は、通知表_全員だけを PDF にしてページ数が人数と同じになることで確かめる。
"""

import datetime as dt
import re
import shutil
import subprocess

import pytest
from openpyxl import Workbook, load_workbook

from grading.export.report_card_pages import (ALL_SHEET, CERT_SHEET, DATA_SHEET, ONE_SHEET, PAGE_ROWS, R_FOOT,
                                              R_PROFILE, R_REM, R_SLOT, R_SUM1, R_TITLE, R_TOTAL, R_ATT,
                                              ROW_HEIGHTS, add_pages, page_height_pt, slot_col)

LONG_NAME = "MUHAMMAD ABDUL RAHMAN BIN ISKANDAR ZULKARNAIN AL-FARUQI"


def _students():
    """(学籍番号, 英字, カタカナ, 学科, 科目[(科目,形態,点数,評定,設定,取得)], 合計(設定,取得,時数,出席,率,不合格,要確認,特記))。"""
    return [
        ("AIBC26001", "TARO YAMADA", "ヤマダ タロウ", "国際ビジネス科",
         [("ビジネス日本語Ⅰ", "講義", 85.04, "A", 2, 2), ("マーケティングⅠ", "講義", 31.1, "E", 2, "－"),
          ("AI演習（実践）Ⅰ", "講義", 70.05, "B", 3, 3)],
         (7, 5, 120, 110.6666666, 110.6666666 / 120, True, 0, "")),
        ("AIBC26002", LONG_NAME, "ムハンマド", "総合ビジネス科",
         [("ビジネス日本語Ⅰ", "講義", "", "F", 2, ""), ("マーケティングⅠ", "講義", 90, "A", 2, 2)],
         (4, 2, 60, 58, 58 / 60, True, 0, "日本語能力強化演習は例外処置")),
        ("AIBC26003", "HANA SATO", "サトウ ハナ", "国際ビジネス科",
         [("ビジネス日本語Ⅰ", "講義", 95, "A", 2, 2)],
         (2, 2, 30, 30, 1, False, 0, "")),
        ("AIBC26004", "JIRO SUZUKI", "スズキ ジロウ", "国際ビジネス科",
         [("ビジネス日本語Ⅰ", "講義", "要確認", "要確認", 2, "要確認")],
         (2, "要確認", 30, 27, 0.9, False, 2, "")),
    ]


def _book(students=None):
    wb = Workbook()
    ws = wb.active
    ws.title = DATA_SHEET
    ws["A1"] = "通知表データ（試し）"
    for k, (sid, en, kana, dept, lines, totals) in enumerate(students or _students()):
        r = 4 + k
        for c, v in zip("ABCD", (sid, en, kana, dept)):
            ws[f"{c}{r}"] = v
        for j in range(1, 14):
            line = lines[j - 1] if j <= len(lines) else ("",) * 6
            for f, v in enumerate(line):
                ws[f"{slot_col(j, f)}{r}"] = v
        for c, v in zip(("CE", "CF", "CG", "CH", "CI", "CJ", "CK", "CL"), totals):
            ws[f"{c}{r}"] = v
    return wb


def _has_calc() -> bool:
    return shutil.which("soffice") is not None and subprocess.run(
        ["dpkg", "-s", "libreoffice-calc"], capture_output=True).returncode == 0


needs_calc = pytest.mark.skipif(not _has_calc(), reason="LibreOffice Calc が無いため式の再計算を確認できない")


def _convert(wb, tmp_path, fmt):
    src = tmp_path / "s.xlsx"
    wb.save(src)
    out = tmp_path / "out"
    subprocess.run(["soffice", f"-env:UserInstallation=file://{tmp_path}/profile", "--headless", "--convert-to",
                    fmt, "--outdir", str(out), str(src)], check=True, capture_output=True, timeout=300)
    return out


def _calc(wb, tmp_path):
    return load_workbook(_convert(wb, tmp_path, "xlsx:Calc MS Excel 2007 XML") / "s.xlsx", data_only=True)


def _grade_rows(ws, top):
    rows = []
    for i in range(12):
        r = top + R_SLOT + i
        rows.append(tuple(ws.cell(r, c).value for c in (1, 3, 4, 5, 6, 7)))
    return rows


def _filled(rows):
    return [r for r in rows if any(v not in (None, "") for v in r)]


# ---------- 式を書いただけで分かること ----------

def test_page_fits_one_a4_portrait_at_100_percent():
    assert len(ROW_HEIGHTS) == PAGE_ROWS
    assert sum(ROW_HEIGHTS) <= page_height_pt() - 20     # LibreOffice・Excel の丸めに 20pt の余裕


def test_all_sheet_has_one_break_between_pages_and_a4_portrait_print_setup():
    wb = add_pages(_book(), students=4)
    ws = wb[ALL_SHEET]
    assert [b.id for b in ws.row_breaks.brk] == [PAGE_ROWS, 2 * PAGE_ROWS, 3 * PAGE_ROWS]
    for name in (ALL_SHEET, ONE_SHEET, CERT_SHEET):
        s = wb[name]
        assert int(s.page_setup.paperSize) == 9 and s.page_setup.orientation == "portrait"
        # 横は1ページ幅に合わせ、縦は改ページで区切る（本物のデータで右端がはみ出して2ページに割れたため）
        assert s.sheet_properties.pageSetUpPr.fitToPage and s.page_setup.fitToWidth == 1 and s.page_setup.fitToHeight == 0
        assert s.print_area
    assert ws.print_area == f"'{ALL_SHEET}'!$A$1:$O${4 * PAGE_ROWS}"


def test_default_is_sixty_pages_that_differ_only_in_the_data_row():
    ws = add_pages(_book())[ALL_SHEET]
    assert len(ws.row_breaks.brk) == 59
    first = [[c.value for c in row] for row in ws.iter_rows(min_row=1, max_row=PAGE_ROWS)]
    last = [[c.value for c in row] for row in ws.iter_rows(min_row=59 * PAGE_ROWS + 1, max_row=60 * PAGE_ROWS)]
    swap = lambda v: re.sub(r"\$63\b", "$4", v) if isinstance(v, str) else v
    assert [[swap(v) for v in row] for row in last] == first


def test_black_rows_and_stamp_are_conditional_formats_on_report_cards_only():
    wb = add_pages(_book(), students=4)

    def formulas(ws):
        return [(str(cf.sqref), rule.formula) for cf in ws.conditional_formatting for rule in cf.rules]
    rules = formulas(wb[ALL_SHEET])
    black = [(ref, f) for ref, f in rules if f and "<>\"A\"" in f[0]]
    assert len(black) == 4
    assert black[0][0] == f"A{1 + R_SLOT}:G{R_SLOT + 12}"
    assert any(f == ['"不合格"'] for _, f in rules)
    assert any(f and "<>\"A\"" in f[0] for _, f in formulas(wb[ONE_SHEET]))
    cert = formulas(wb[CERT_SHEET])
    assert not any(f and ("<>\"A\"" in f[0] or f == ['"不合格"']) for _, f in cert)
    assert any(f == ['"要確認"'] for _, f in cert)


# ---------- LibreOffice で再計算して分かること ----------

@needs_calc
def test_each_page_shows_its_student(tmp_path):
    ws = _calc(add_pages(_book(), students=4), tmp_path)[ALL_SHEET]
    tops = [1 + k * PAGE_ROWS for k in range(4)]
    p = lambda top, i, c: ws.cell(top + R_PROFILE + i, c).value
    assert [(p(t, 0, 2), p(t, 0, 5), p(t, 1, 5), p(t, 0, 13)) for t in tops] == [
        ("国際ビジネス科", "ヤマダ タロウ", "TARO YAMADA", "AIBC26001"),
        ("総合ビジネス科", "ムハンマド", LONG_NAME, "AIBC26002"),
        ("国際ビジネス科", "サトウ ハナ", "HANA SATO", "AIBC26003"),
        ("国際ビジネス科", "スズキ ジロウ", "JIRO SUZUKI", "AIBC26004")]
    assert ws.cell(1 + R_TITLE, 3).value == "通　知　表"
    assert ws.cell(1, 11).value == "この書類は成績証明書ではありません"
    today = dt.date.today()
    assert ws.cell(1 + R_PROFILE + 1, 13).value.date() == today
    assert ws.cell(1 + R_FOOT, 9).value == f"発行年月日　{today.year}年{today.month}月{today.day}日"
    assert ws.cell(1 + R_PROFILE + 2, 2).value == "2026年度 前期"
    assert ws.cell(1 + R_FOOT, 1).value == "記載の期間における出席及び成績は、上記のとおりです。"


@needs_calc
def test_subjects_are_packed_and_shown_in_display_order(tmp_path):
    ws = _calc(add_pages(_book(), students=4), tmp_path)[ALL_SHEET]
    # 授業科目・点数・評定・形態・設定・取得
    assert _filled(_grade_rows(ws, 1)) == [
        ("ビジネス日本語Ⅰ", 85, "A", "講義", 2, 2), ("マーケティングⅠ", 31.1, "E", "講義", 2, "－"),
        ("AI演習（実践）Ⅰ", 70.1, "B", "講義", 3, 3)]
    assert len(_filled(_grade_rows(ws, 1)[:3])) == 3
    assert ws.cell(1 + R_SLOT, 3).number_format == "General"
    exception = _filled(_grade_rows(ws, 1 + PAGE_ROWS))[0]
    assert exception == ("ビジネス日本語Ⅰ", None, "F", "講義", 2, None)
    assert [ws.cell(1 + R_SUM1, c).value for c in (1, 6, 7)] == ["前期計", 7, 5]
    assert [ws.cell(1 + R_TOTAL, c).value for c in (1, 6, 7)] == ["1年 合計取得単位数", 7, 5]
    assert [ws.cell(1 + R_TOTAL, c).value for c in (9, 14, 15)] == ["2年 合計取得単位数", None, None]


@needs_calc
def test_attendance_remarks_stamp_and_check_note(tmp_path):
    ws = _calc(add_pages(_book(), students=4), tmp_path)[ALL_SHEET]
    top1, top2, top3, top4 = (1 + k * PAGE_ROWS for k in range(4))
    att = [tuple(ws.cell(top1 + R_ATT + i, c).value for c in (1, 2, 3, 5)) for i in range(3)]
    assert att[0][:3] == ("前期", 120, 110.7)
    assert att[0][3] == pytest.approx(110.6666666 / 120)
    assert att[1] == ("後期", None, None, None)
    assert att[2][:2] == ("年間合計", 120)
    assert ws.cell(top1 + R_ATT, 3).number_format == "General" and ws.cell(top1 + R_ATT, 5).number_format == "0.0%"
    assert [ws.cell(top1 + R_ATT + i, 10).value for i in range(3)] == [None] * 3          # 2年は空欄
    assert ws.cell(top1 + R_REM, 1).value == "該当なし"
    assert ws.cell(top2 + R_REM, 1).value == "日本語能力強化演習は例外処置"
    stamps = [ws.cell(t + R_TITLE, 13).value for t in (top1, top2, top3, top4)]
    assert stamps == ["不合格", "不合格", None, None]
    notes = [ws.cell(t + R_TITLE, 1).value for t in (top1, top2, top3, top4)]
    assert notes == [None, None, None, "要確認あり（2件）"]
    assert _filled(_grade_rows(ws, top4))[0][1:3] == ("要確認", "要確認")


@needs_calc
def test_single_pages_follow_the_entered_id(tmp_path):
    wb = add_pages(_book(), students=4)
    wb[ONE_SHEET]["B1"] = "AIBC26002"
    wb[CERT_SHEET]["B1"] = "AIBC26001"
    v = _calc(wb, tmp_path)
    one, cert = v[ONE_SHEET], v[CERT_SHEET]
    top = 3
    assert one.cell(top + R_PROFILE + 1, 5).value == LONG_NAME
    assert one.cell(top + R_TITLE, 3).value == "通　知　表" and one.cell(top + R_TITLE, 13).value == "不合格"
    assert one.cell(top, 11).value == "この書類は成績証明書ではありません"
    assert one.cell(top + R_REM, 1).value == "日本語能力強化演習は例外処置"
    assert cert.cell(top + R_TITLE, 3).value == "成　績　証　明　書"
    assert cert.cell(top + R_TITLE, 13).value is None
    assert cert.cell(top, 11).value is None
    assert cert.cell(top + R_PROFILE, 13).value == "AIBC26001"
    assert _filled(_grade_rows(cert, top))[1] == ("マーケティングⅠ", 31.1, "E", "講義", 2, "－")
    assert [cert.cell(top + R_SUM1, c).value for c in (6, 7)] == [7, 5]
    assert cert.cell(top + R_ATT, 2).value == 120


@needs_calc
def test_unknown_id_shows_check_in_red_and_blank_id_shows_nothing(tmp_path):
    wb = add_pages(_book(), students=4)
    wb[ONE_SHEET]["B1"] = "AIBC99999"
    v = _calc(wb, tmp_path)
    one, cert = v[ONE_SHEET], v[CERT_SHEET]
    top = 3
    assert one.cell(top + R_PROFILE, 13).value == "要確認"
    assert one.cell(top + R_PROFILE + 1, 5).value == "要確認"
    assert one.cell(top + R_TITLE, 1).value.startswith("要確認")
    assert one.cell(top + R_TITLE, 1).font.color.rgb.endswith("C00000")
    assert _filled(_grade_rows(one, top)) == []
    assert cert.cell(top + R_PROFILE, 13).value is None          # 未入力は空欄（要確認にしない）


@needs_calc
def test_all_sheet_prints_one_page_per_student(tmp_path):
    """通知表_全員だけを PDF にして、ページ数＝人数（60人ぶんのシート）。"""
    tmp_calc, tmp_pdf = tmp_path / "calc", tmp_path / "pdf"
    tmp_calc.mkdir(), tmp_pdf.mkdir()
    # LibreOffice は隠したシートも PDF にするので、再計算した値の本から通知表_全員だけを残して PDF にする
    wb = _calc(add_pages(_book()), tmp_calc)
    for name in wb.sheetnames:
        if name != ALL_SHEET:
            del wb[name]
    pdf = _convert(wb, tmp_pdf, "pdf") / "s.pdf"
    try:
        from pypdf import PdfReader
        pages = len(PdfReader(pdf).pages)
    except ImportError:
        pages = len(re.findall(rb"/Type\s*/Page\b", pdf.read_bytes()))
    assert pages == 60


@needs_calc
def test_thirteenth_subject_is_flagged_because_the_page_has_twelve_rows(tmp_path):
    many = [(f"科目{i}", "講義", 80, "B", 1, 1) for i in range(13)]
    students = [("AIBC26001", "TARO", "タロウ", "国際ビジネス科", many, (13, 13, 30, 30, 1, False, 0, ""))]
    ws = _calc(add_pages(_book(students), students=1), tmp_path)[ALL_SHEET]
    assert ws.cell(1 + R_TITLE, 1).value.startswith("要確認")
