import re

from openpyxl import Workbook

from grading.term2.first_term import SHEET as FIRST_TERM_SHEET
from grading.term2.layout import (CHECK_SHEET, EXCEPTION_SHEET, FIRST_ROW, INTAKE_SHEET, REMARKS_SHEET, ROSTER_SHEET,
                                  SUMMARY_SHEET, gradebook_cells)
from grading.term2.report_data import (REPORT_SHEET, TOTALS, _col, exception_cells, remarks_cells,
                                       report_data_cells)

from .test_term2_layout import _calc, _sheet, needs_calc

EX = "COUNTIFS(例外!$A$4:$A$53,$A4,例外!$B$4:$B$53,CE4)>0"


def _num(col: str) -> int:
    n = 0
    for ch in col:
        n = n * 26 + ord(ch) - ord("A") + 1
    return n


def _at(rows, addr):
    m = re.fullmatch(r"([A-Z]+)(\d+)", addr)
    return rows[int(m[2]) - 1][_num(m[1]) - 1]


def _shift(v, r):
    """行番号 r の相対参照を # に置き換える（行によらず同じ式かを比べる）。"""
    return re.sub(rf"(?<=[A-Z]){r}\b", "#", str(v))


# ---------- 通知表データ ----------

def test_report_data_title_and_headers():
    rows = report_data_cells().tabs[REPORT_SHEET]
    assert rows[0][0] == "通知表データ（通知表_印刷用が読む。1人1行。直さない）"
    assert len(rows) == 63 and all(len(r) == _num("FV") for r in rows)
    head = rows[2]
    assert head[:4] == ["学籍番号", "氏名", "カタカナ", "学科"]
    first = ('="前期"&(INT((COLUMN()-5)/6)+1)&"_"&CHOOSE(MOD(COLUMN()-5,6)+1,'
             '"科目","形態","点数","評定","設定","取得")')
    assert head[_num("E") - 1:_num("CD")] == [first] * 78                  # E3 を E3:CD3 へ写したもの
    assert head[_num("CE") - 1:_num("FD")] == [first.replace("前期", "後期").replace("-5", "-83")] * 78
    assert head[_num("FE") - 1:] == TOTALS
    assert TOTALS[0] == "前期設定計" and TOTALS[-1] == "特記事項" and len(TOTALS) == 18


def test_report_data_row4_formulas():
    rows = report_data_cells().tabs[REPORT_SHEET]
    at = lambda a: _at(rows, a)
    assert at("A4") == '=IFERROR(INDEX(名簿!$A$4:$A$203,ROW()-3),"")'
    for c, src in zip("BCD", "BCD"):
        assert at(f"{c}4") == f'=IF($A4="","",INDEX(名簿!${src}$4:${src}$203,ROW()-3))'
    assert at("E4") == ('=IF($A4="","",IFERROR(INDEX(FILTER(前期成績!$D$4:$I$723,前期成績!$A$4:$A$723=$A4),'
                        'INT((COLUMN()-5)/6)+1,MOD(COLUMN()-5,6)+1),""))')
    assert at("CE4") == ('=IF($A4="","",IFERROR(INDEX(FILTER(集計!$A$4:$A$867,集計!$B$4:$B$867=$A4),'
                         'INT((COLUMN()-83)/6)+1),""))')
    subject = lambda c: (f'=IF(CE4="","",IFERROR(INDEX(FILTER(科目!${c}$4:${c}$100,科目!$A$4:$A$100=$D4,'
                         f'科目!$B$4:$B$100=CE4),1),""))')
    assert at("CF4") == subject("E") and at("CI4") == subject("D")
    assert at("CG4") == (f'=IF(CE4="","",IF({EX},"",IFERROR(INDEX(FILTER(集計!$F$4:$F$867,集計!$B$4:$B$867=$A4),'
                         f'INT((COLUMN()-83)/6)+1),"")))')
    assert at("CH4") == (f'=IF(CE4="","",IF({EX},IFERROR(INDEX(FILTER(例外!$C$4:$C$53,例外!$A$4:$A$53=$A4,'
                         f'例外!$B$4:$B$53=CE4),1),""),IFERROR(INDEX(FILTER(集計!$H$4:$H$867,集計!$B$4:$B$867=$A4),'
                         f'INT((COLUMN()-83)/6)+1),"")))')
    assert at("CJ4") == f'=IF(OR(CE4="",CH4="",CH4="E",CH4="F",CH4="要確認",{EX}),"",CI4)'
    first = lambda c: f'=IF($A4="","",IFERROR(INDEX(前期成績!${c}$4:${c}$63,MATCH($A4,前期成績!$K$4:$K$63,0)),""))'
    assert [at(a) for a in ("FE4", "FF4", "FK4", "FL4")] == [first(c) for c in "LMNO"]
    assert at("FG4") == '=IF($A4="","",SUM(CI4,CO4,CU4,DA4,DG4,DM4,DS4,DY4,EE4,EK4,EQ4,EW4,FC4))'
    assert at("FH4") == '=IF($A4="","",SUM(CJ4,CP4,CV4,DB4,DH4,DN4,DT4,DZ4,EF4,EL4,ER4,EX4,FD4))'
    assert at("FI4") == '=IF($A4="","",N(FE4)+FG4)' and at("FJ4") == '=IF($A4="","",N(FF4)+FH4)'
    ik = lambda c: f"'出席の受け口'!${c}$4:${c}$843"
    assert at("FM4") == f'=IF($A4="","",SUMIFS({ik("I")},{ik("A")},$A4))'
    assert at("FN4") == f'=IF($A4="","",FM4-SUMIFS({ik("J")},{ik("A")},$A4)-SUMIFS({ik("K")},{ik("A")},$A4)/3)'
    assert at("FO4") == '=IF($A4="","",N(FK4)+N(FM4))' and at("FP4") == '=IF($A4="","",N(FL4)+N(FN4))'
    assert [at(a) for a in ("FQ4", "FR4", "FS4")] == ['=IF(N(FK4)=0,"",FL4/FK4)', '=IF(N(FM4)=0,"",FN4/FM4)',
                                                      '=IF(N(FO4)=0,"",FP4/FO4)']
    assert at("FT4") == ('=IF($A4="","",OR(COUNTIF($E4:$FD4,"E")+COUNTIF($E4:$FD4,"F")>0,'
                         'COUNTIF(例外!$A$4:$A$53,$A4)>0))')
    assert at("FU4") == '=IF($A4="","",OR(COUNTIF($CE4:$FD4,"要確認")>0,確認!$B$3>0,確認!$D$3>0))'
    assert at("FV4") == ('=IF($A4="","",IFERROR(INDEX(FILTER(特記事項!$B$4:$B$103,特記事項!$A$4:$A$103=$A4),1),'
                         '"該当なし"))')


def test_report_data_slots_are_copies_across_columns():
    row = report_data_cells().tabs[REPORT_SHEET][FIRST_ROW - 1]
    assert len(set(row[_num("E") - 1:_num("CD")])) == 1                    # 前期: E4 を E4:CD4 へ写したもの
    # 後期: CE4:CJ4 を 6列ずつ右へ写したもの（相対参照の列が6つずつずれる）
    back = lambda v, k: re.sub(r"(?<![$A-Z])([A-Z]{1,3})(?=4\b)",
                               lambda m: _col(_num(m[1]) - 6 * k) if _num(m[1]) >= _num("CE") else m[1], v)
    base = row[_num("CE") - 1:_num("CJ")]
    for k in range(13):
        block = row[_num("CE") - 1 + 6 * k:_num("CJ") + 6 * k]
        assert [back(v, k) for v in block] == base, k
    assert block[0].startswith('=IF($A4="","",') and "FD4" not in back(block[5], 12)


def test_report_data_rows_are_copies_down():
    rows = report_data_cells().tabs[REPORT_SHEET][FIRST_ROW - 1:]
    assert len(rows) == 60
    assert len({tuple(_shift(v, FIRST_ROW + k) for v in row) for k, row in enumerate(rows)}) == 1


def test_report_data_ranges_follow_gradebook_size():
    """集計・出席の受け口の下端は gradebook_cells の確認シートと同じ数え方で決まる。"""
    rows = report_data_cells(subjects=2, slots=3, students=5, first_term_lines=10).tabs[REPORT_SHEET]
    check = gradebook_cells([("国際", "科目A"), ("総合", "科目A")], slots=3).tabs[CHECK_SHEET]
    assert len(rows) == FIRST_ROW - 1 + 5
    assert "集計!$A$4:$A$11" in _at(rows, "CE4") and "集計!$I$4:$I$11" in check[2][3]
    assert "'出席の受け口'!$I$4:$I$9" in _at(rows, "FM4") and "'出席の受け口'!$L$4:$L$9" in check[2][3]
    assert "前期成績!$D$4:$I$13" in _at(rows, "E4") and "前期成績!$K$4:$K$8" in _at(rows, "FE4")


# ---------- 例外・特記事項 ----------

def test_exception_sheet_layout():
    book = exception_cells()
    rows = book.tabs[EXCEPTION_SHEET]
    assert rows[0][0] == "例外" and rows[1][0].startswith("1行書くと、その科目が通知表で黒塗り")
    assert rows[2] == ["学籍番号", "科目", "表示する評定", "理由", "確認"]
    assert rows[3][4] == ('=IF(AND($A4="",$B4=""),"",IF(COUNTIFS(集計!$B$4:$B$867,$A4,集計!$A$4:$A$867,$B4)=0,'
                          '"学籍番号または科目が見つからない",IF($C4="","表示する評定が無い","")))')
    body = rows[FIRST_ROW - 1:]
    assert len(body) == 50 and all(r[:4] == [None] * 4 for r in body)
    assert len({_shift(r[4], FIRST_ROW + k) for k, r in enumerate(body)}) == 1
    assert book.yellow == {EXCEPTION_SHEET: ["A4:D53"]}
    assert book.red == {EXCEPTION_SHEET: [("E4:E53", '=$E4<>""')]}


def test_remarks_sheet_layout():
    book = remarks_cells()
    rows = book.tabs[REMARKS_SHEET]
    assert rows[0][0] == "特記事項" and rows[1][0] == "書いた学生だけ、その文が通知表に出る。書かなければ「該当なし」"
    assert rows[2] == ["学籍番号", "文", "確認"]
    assert rows[3][2] == ('=IF(AND($A4="",$B4=""),"",IF(COUNTIF(名簿!$A$4:$A$203,$A4)=0,"学籍番号が名簿に無い",'
                          'IF($B4="","文が無い","")))')
    body = rows[FIRST_ROW - 1:]
    assert len(body) == 100 and len({_shift(r[2], FIRST_ROW + k) for k, r in enumerate(body)}) == 1
    assert book.yellow == {REMARKS_SHEET: ["A4:B103"]}
    assert book.red == {REMARKS_SHEET: [("C4:C103", '=$C4<>""')]}


def test_check_sheet_counts_and_lists_exceptions_and_remarks():
    book = gradebook_cells([("国際", f"科目{k}") for k in range(24)])
    check = book.tabs[CHECK_SHEET]
    ok = "AND(COUNT(設定!$F$4:$F$7)=4,設定!$F$4>設定!$F$5,設定!$F$5>設定!$F$6,設定!$F$6>設定!$F$7,設定!$F$7>0)"
    assert check[2][3] == (f'=COUNTIF(集計!$I$4:$I$867,"?*")+COUNTIF(\'出席の受け口\'!$L$4:$L$843,"?*")'
                           f'+COUNTIF(例外!$E$4:$E$53,"?*")+COUNTIF(特記事項!$C$4:$C$103,"?*")+IF({ok},0,1)')
    assert check[4][10:13] == ["学籍番号", "科目", "何がおかしいか（例外・特記事項）"]
    assert check[5][10] == ('=IFERROR(FILTER(VSTACK({例外!$A$4:$A$53,例外!$B$4:$B$53,例外!$E$4:$E$53},'
                            '{特記事項!$A$4:$A$103,IF(特記事項!$A$4:$A$103="","","（特記事項）"),特記事項!$C$4:$C$103}),'
                            'VSTACK(例外!$E$4:$E$53,特記事項!$C$4:$C$103)<>""),"")')
    assert ("K6:M1000", '=$M6<>""') in book.red[CHECK_SHEET]
    assert check[4][:9] == ["未提出の科目", None, "学籍番号", "科目", "何がおかしいか（成績）", None,
                            "学籍番号", "科目", "何がおかしいか（出席）"]               # 前からの欄はそのまま


# ---------- LibreOffice で再計算 ----------

@needs_calc
def test_exception_check_flags_unknown_pair_and_missing_grade(tmp_path):
    wb = Workbook()
    wb.remove(wb.active)
    summary = wb.create_sheet(SUMMARY_SHEET)
    summary.cell(FIRST_ROW, 1, "国際　科目A")                               # 見出しの行
    for k, sid in enumerate(["K01", "K02", "K03"], FIRST_ROW + 1):
        summary.cell(k, 1, "科目A"), summary.cell(k, 2, sid)
    cases = [("K01", "科目A", "F", "理由"), ("K01", "科目B", "F", None), ("K09", "科目A", "F", None),
             ("K02", "科目A", None, None), (None, None, None, None), (None, None, "F", "理由だけ")]
    ws = _sheet(wb, EXCEPTION_SHEET, exception_cells(subjects=1, slots=3).tabs[EXCEPTION_SHEET])
    for k, case in enumerate(cases):
        for c, v in enumerate(case, 1):
            ws.cell(FIRST_ROW + k, c, v)
    got = _calc(wb, tmp_path)[EXCEPTION_SHEET]
    lost = "学籍番号または科目が見つからない"
    assert [got.cell(FIRST_ROW + k, 5).value for k in range(len(cases))] == [
        None, lost, lost, "表示する評定が無い", None, None]


@needs_calc
def test_remarks_check_flags_unknown_id_and_missing_text(tmp_path):
    wb = Workbook()
    wb.remove(wb.active)
    roster = wb.create_sheet(ROSTER_SHEET)
    for k, sid in enumerate(["K01", "K02"], FIRST_ROW):
        roster.cell(k, 1, sid)
    cases = [("K01", "皆勤"), ("X99", "皆勤"), ("K02", None), (None, None)]
    ws = _sheet(wb, REMARKS_SHEET, remarks_cells().tabs[REMARKS_SHEET])
    for k, case in enumerate(cases):
        for c, v in enumerate(case, 1):
            ws.cell(FIRST_ROW + k, c, v)
    got = _calc(wb, tmp_path)[REMARKS_SHEET]
    assert [got.cell(FIRST_ROW + k, 3).value for k in range(len(cases))] == [
        None, "学籍番号が名簿に無い", "文が無い", None]


@needs_calc
def test_report_data_totals_rates_and_flags(tmp_path):
    """FILTER を使わない欄（計・出席・出席率・不合格・要確認）を、読み元を値で置いて再計算する。"""
    rows = report_data_cells(subjects=1, slots=3, students=3, first_term_lines=6).tabs[REPORT_SHEET]
    wb = Workbook()
    wb.remove(wb.active)
    fe, fu = _num("FE"), _num("FU")
    ws = _sheet(wb, REPORT_SHEET, rows, keep=lambda c: fe <= c <= fu)
    ce = _num("CE")
    # 学生ごと: 学籍番号, 前期の評定, 後期の科目（評定, 設定, 取得）, 例外
    students = [("K01", "A", [("B", 2, 2), ("C", 1, 1)]),
                ("K02", "E", [("要確認", 2, None)]),
                ("K03", "B", [("F", 2, None)])]
    for k, (sid, first_grade, second) in enumerate(students):
        r = FIRST_ROW + k
        ws.cell(r, 1, sid), ws.cell(r, _num("H"), first_grade)
        for j, (g, credits, earned) in enumerate(second):
            ws.cell(r, ce + 6 * j, f"科目{j}"), ws.cell(r, ce + 6 * j + 3, g)
            ws.cell(r, ce + 6 * j + 4, credits), ws.cell(r, ce + 6 * j + 5, earned)
    first = wb.create_sheet(FIRST_TERM_SHEET)
    for k, (sid, *vals) in enumerate([("K01", 10, 9, 60, 57), ("K02", 10, 7, 60, 50), ("K03", 8, 8, 0, 0)]):
        for c, v in enumerate([sid, *vals], _num("K")):
            first.cell(FIRST_ROW + k, c, v)
    intake = wb.create_sheet(INTAKE_SHEET)
    for k, (sid, n, a, late) in enumerate([("K01", 30, 2, 3), ("K01", 15, 0, 0), ("K02", 30, 0, 0)]):
        for c, v in zip((1, 9, 10, 11), (sid, n, a, late)):
            intake.cell(FIRST_ROW + k, c, v)
    wb.create_sheet(EXCEPTION_SHEET).cell(FIRST_ROW, 1, "K01")              # K01 に例外が1つある
    chk = wb.create_sheet(CHECK_SHEET)
    chk["B3"], chk["D3"] = 0, 0
    got = _calc(wb, tmp_path)[REPORT_SHEET]
    val = lambda col, k: got.cell(FIRST_ROW + k, _num(col)).value
    assert [val(c, 0) for c in ("FE", "FF", "FG", "FH", "FI", "FJ")] == [10, 9, 3, 3, 13, 12]
    assert [val(c, 0) for c in ("FK", "FL", "FM", "FN", "FO", "FP")] == [60, 57, 45, 42, 105, 99]
    assert abs(val("FQ", 0) - 57 / 60) < 1e-9 and abs(val("FS", 0) - 99 / 105) < 1e-9
    assert val("FR", 2) is None and val("FQ", 2) is None                    # 授業時数が0なら出席率は空欄
    assert [val("FT", k) for k in range(3)] == [True, True, True]          # 例外・前期のE・後期のF
    assert [val("FU", k) for k in range(3)] == [False, True, False]        # 後期に要確認
