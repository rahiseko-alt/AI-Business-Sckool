from fractions import Fraction

import pytest
from openpyxl import Workbook, load_workbook

from grading.export.fill import FilledValue, fill_template


def _blank(path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "2026年度　国際ビジネス科（AI計算版）"
    ws["B3"] = "学籍番号"
    ws["E3"], ws["E4"], ws["F4"], ws["G4"] = "マーケティング", "出席", "授業態度", "筆記テスト"
    ws["H3"], ws["H4"], ws["I4"] = "AI演習（実践）", "出席", "授業態度"
    ws.merge_cells("E3:G3")
    ws["B6"], ws["B7"] = "AIBC26001", "AIBC26002"
    wb.save(path)
    return path


def _v(sid, subject, item, value, ev=("出席簿!F6",)):
    return FilledValue(sid, subject, item, Fraction(value), "10 − 2 × 遅刻1回 = 8", tuple(ev))


def test_values_go_to_the_cell_found_by_subject_and_item(tmp_path):
    out = fill_template(_blank(tmp_path / "b.xlsx"), tmp_path / "o.xlsx",
                        [_v("AIBC26001", "マーケティング", "授業態度", 8), _v("AIBC26002", "AI演習", "授業態度", 10, ())])
    ws = load_workbook(out).worksheets[0]
    assert ws["F6"].value == 8 and ws["I7"].value == 10
    assert ws["E6"].value is None


def test_every_value_is_listed_with_its_formula_and_source(tmp_path):
    out = fill_template(_blank(tmp_path / "b.xlsx"), tmp_path / "o.xlsx", [_v("AIBC26001", "マーケティング", "授業態度", 8)])
    rows = list(load_workbook(out)["根拠"].iter_rows(min_row=2, values_only=True))
    assert rows == [("AIBC26001", "マーケティング", "授業態度", "F6", 8, "8", "10 − 2 × 遅刻1回 = 8", "出席簿!F6")]


def test_fractions_are_kept_exactly_in_the_evidence_sheet(tmp_path):
    out = fill_template(_blank(tmp_path / "b.xlsx"), tmp_path / "o.xlsx", [_v("AIBC26001", "マーケティング", "筆記テスト", Fraction(1, 3))])
    row = next(load_workbook(out)["根拠"].iter_rows(min_row=2, values_only=True))
    assert row[5] == "1/3"


@pytest.mark.parametrize("sid, subject, item", [("AIBC26999", "マーケティング", "授業態度"),
                                                ("AIBC26001", "経営学", "授業態度"),
                                                ("AIBC26001", "マーケティング", "発表")])
def test_unknown_student_subject_or_item_stops_instead_of_guessing(tmp_path, sid, subject, item):
    with pytest.raises(KeyError):
        fill_template(_blank(tmp_path / "b.xlsx"), tmp_path / "o.xlsx", [_v(sid, subject, item, 8)])


def test_writing_the_same_cell_twice_is_refused(tmp_path):
    with pytest.raises(ValueError):
        fill_template(_blank(tmp_path / "b.xlsx"), tmp_path / "o.xlsx",
                      [_v("AIBC26001", "マーケティング", "授業態度", 8), _v("AIBC26001", "マーケティング", "授業態度", 6)])
