"""使い方シート（前期の成績表を AI なしで直すための手順書き）。"""

from openpyxl import Workbook

from grading.export.report_card_guide import GUIDE_SHEET, SHEETS_NAMED, add_guide


def _book():
    wb = Workbook()
    wb.active.title = "AI計算版_国際"
    return wb


def test_guide_is_first_and_opens_first():
    wb = _book()
    add_guide(wb)
    assert wb.sheetnames[0] == GUIDE_SHEET
    assert wb.active.title == GUIDE_SHEET


def test_guide_text_covers_each_step():
    wb = _book()
    add_guide(wb)
    text = "\n".join(str(c.value) for row in wb[GUIDE_SHEET].iter_rows() for c in row if c.value)
    for word in ("点数", "出席", "例外", "特記事項", "確認", "B2", "通知表_1人", "成績証明書_1人",
                 "通知表_全員", "保存", "=", "黄色"):
        assert word in text, word


def test_guide_has_no_formulas():
    wb = _book()
    add_guide(wb)
    for row in wb[GUIDE_SHEET].iter_rows():
        for c in row:
            assert not (isinstance(c.value, str) and c.value.startswith("=")), c.coordinate


def test_named_sheets_are_the_built_ones():
    from grading.export import report_card_pages as p
    for name in (p.ALL_SHEET, p.ONE_SHEET, p.CERT_SHEET):
        assert name in SHEETS_NAMED
